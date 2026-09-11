#!/usr/bin/env python3
"""Internal Halex (Fred 9 Sep 2026) FOB → landed compare. Do not send to the mill."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT_DIR = Path(__file__).resolve().parent
COMPARE_XLSX = OUT_DIR / "Halex_Fred_quote_COMPARE.xlsx"
COMPARE_CSV = OUT_DIR / "Halex_Fred_quote_COMPARE.csv"

DUTY = 1.43
OCEAN = 10_000.0
LF_PER_BOX = 400  # 100 pcs × 4 ft

HEADER = PatternFill("solid", fgColor="C00000")
WHITE_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
GREEN = PatternFill("solid", fgColor="C6EFCE")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="center")
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)

# Fred sheet 20260909 — FOB Shanghai USD/box. MOQ = 40HQ mixed SKUs.
# US public check: Home Depot Halex 100-pack (same mill brand, retail box).
SKUS = [
    {
        "sku": "SGR-120",
        "maps_to": "HL-120 / 7/8\" wood",
        "description": "REGULAR WIDE WOOD 3/4\" RS",
        "wood": "All poplar L 1220mm x W 21.6mm x H 6.5mm",
        "width_in": '7/8"',
        "nails": 'Wood nail 13GA x 3/4" ring shank, tumbled, 9 pcs/strip',
        "pins": "Tumbled pins, 52 pcs/strip, ~60°",
        "carb": "Yes",
        "pcs_box": 100,
        "fob": 10.74,
        "boxes_40hq": 1740,
        "cbm_40hq": 73.0,
        "boxes_pallet": 40,
        "cbm_pallet": 1.72,
        "gw_pallet_kg": 418,
        "load_note": "1740 (36 pallets + hand stack 300 boxes)",
        "us_sku": "HD-320-P-8",
        "us_box": 28.65,
        "us_note": "Home Depot wood 100-pack (listed 1 in W; closest public box)",
        "verdict": "Beats US box. Take for wood subfloor.",
    },
    {
        "sku": "SGR-340",
        "maps_to": "HL-340 / 1\" concrete",
        "description": 'EXTRA WIDE CONC 11/16"',
        "wood": "All poplar L 1220mm x W 23.5mm x H 6.5mm",
        "width_in": 'extra-wide ~15/16"',
        "nails": 'Concrete nail 12GA x 11/16", tumbled, 9 pcs/strip',
        "pins": "Tumbled pins, 52 pcs/strip, ~60°",
        "carb": "Yes",
        "pcs_box": 100,
        "fob": 11.97,
        "boxes_40hq": 1540,
        "cbm_40hq": 71.7,
        "boxes_pallet": 36,
        "cbm_pallet": 1.70,
        "gw_pallet_kg": 450,
        "load_note": "1540 (36 pallets + hand stack 244 boxes)",
        "us_sku": "HD-340-P-8",
        "us_box": 37.47,
        "us_note": "Home Depot concrete 100-pack",
        "verdict": "Beats US box. Dual SGR-360 is only $0.10 FOB more — prefer dual unless we want concrete-only.",
    },
    {
        "sku": "SGR-360",
        "maps_to": "dual-purpose extra-wide",
        "description": 'EXTRA WIDE Dual-Purpose 11/16"',
        "wood": "All poplar L 1220mm x W 23.5mm x H 6.5mm",
        "width_in": 'extra-wide ~15/16"',
        "nails": 'Dual-purpose nail 12GA x 11/16", tumbled, 9 pcs/strip',
        "pins": "Tumbled pins, 52 pcs/strip, ~60°",
        "carb": "Yes",
        "pcs_box": 100,
        "fob": 12.07,
        "boxes_40hq": 1540,
        "cbm_40hq": 71.7,
        "boxes_pallet": 36,
        "cbm_pallet": 1.70,
        "gw_pallet_kg": 450,
        "load_note": "1540 (36 pallets + hand stack 244 boxes)",
        "us_sku": "HD-340-P-8",
        "us_box": 37.47,
        "us_note": "No dual SKU at HD; check vs concrete 100-pack",
        "verdict": "Beats US box. Stock this instead of SGR-340 — covers wood + slab for $0.10/box.",
    },
]


def landed(row: dict) -> dict:
    fob = row["fob"]
    boxes = row["boxes_40hq"]
    duty_on_fob = round(fob * DUTY, 4)
    ocean_box = round(OCEAN / boxes, 4)
    landed_box = round(duty_on_fob + ocean_box, 4)
    fob_lf = round(fob / LF_PER_BOX, 5)
    landed_lf = round(landed_box / LF_PER_BOX, 5)
    us_lf = round(row["us_box"] / LF_PER_BOX, 5)
    beat = round(row["us_box"] - landed_box, 2)
    fob_40hq = round(fob * boxes, 2)
    return {
        **row,
        "lf_box": LF_PER_BOX,
        "duty_stack": DUTY,
        "ocean_container": OCEAN,
        "fob_x_duty": duty_on_fob,
        "ocean_per_box": ocean_box,
        "landed_box": landed_box,
        "fob_per_lf": fob_lf,
        "landed_per_lf": landed_lf,
        "us_per_lf": us_lf,
        "beat_vs_us": beat,
        "fob_40hq": fob_40hq,
        "moq": "40HQ (SKUs mixed)",
        "lead": "2–3 weeks after deposit",
        "incoterms": "FOB Shanghai",
        "loading": "Palletized + hand stack to fill",
    }


def mix_row() -> dict:
    """Suggested first 40HQ: 400 wood + 1140 dual (same box count as extra-wide fill)."""
    n120, n360 = 400, 1140
    r120 = next(s for s in SKUS if s["sku"] == "SGR-120")
    r360 = next(s for s in SKUS if s["sku"] == "SGR-360")
    cbm120 = r120["cbm_40hq"] / r120["boxes_40hq"]
    cbm360 = r360["cbm_40hq"] / r360["boxes_40hq"]
    boxes = n120 + n360
    fob_total = r120["fob"] * n120 + r360["fob"] * n360
    fob_avg = fob_total / boxes
    cbm = n120 * cbm120 + n360 * cbm360
    duty_on_fob = round(fob_avg * DUTY, 4)
    ocean_box = round(OCEAN / boxes, 4)
    landed_box = round(duty_on_fob + ocean_box, 4)
    us_blend = (r120["us_box"] * n120 + r360["us_box"] * n360) / boxes
    return {
        "sku": "MIX 400×SGR-120 + 1140×SGR-360",
        "maps_to": "first 40HQ sketch",
        "description": "Wood + dual mixed load (not on Fred sheet — packing math)",
        "wood": "see component SKUs",
        "width_in": "mix",
        "nails": "see component SKUs",
        "pins": "~60°",
        "carb": "Yes",
        "pcs_box": 100,
        "fob": round(fob_avg, 4),
        "boxes_40hq": boxes,
        "cbm_40hq": round(cbm, 2),
        "boxes_pallet": "",
        "cbm_pallet": "",
        "gw_pallet_kg": "",
        "load_note": f"{n120} SGR-120 + {n360} SGR-360; confirm CBM with Fred",
        "us_sku": "blend HD-320 / HD-340",
        "us_box": round(us_blend, 2),
        "us_note": "Weighted HD box vs mix",
        "verdict": "Still beats blended US box. Confirm cube with Fred before PI.",
        "lf_box": LF_PER_BOX,
        "duty_stack": DUTY,
        "ocean_container": OCEAN,
        "fob_x_duty": duty_on_fob,
        "ocean_per_box": ocean_box,
        "landed_box": landed_box,
        "fob_per_lf": round(fob_avg / LF_PER_BOX, 5),
        "landed_per_lf": round(landed_box / LF_PER_BOX, 5),
        "us_per_lf": round(us_blend / LF_PER_BOX, 5),
        "beat_vs_us": round(us_blend - landed_box, 2),
        "fob_40hq": round(fob_total, 2),
        "moq": "40HQ (SKUs mixed)",
        "lead": "2–3 weeks after deposit",
        "incoterms": "FOB Shanghai",
        "loading": "Palletized + hand stack to fill",
    }


COLS = [
    ("sku", "SKU", 28),
    ("maps_to", "Maps to", 26),
    ("description", "Fred description", 32),
    ("width_in", "Width", 16),
    ("nails", "Nails", 42),
    ("pins", "Pins", 28),
    ("carb", "CARB/EPA", 12),
    ("pcs_box", "pcs/box", 10),
    ("lf_box", "LF/box", 10),
    ("incoterms", "Incoterms", 14),
    ("fob", "FOB USD/box", 14),
    ("fob_per_lf", "FOB $/LF", 12),
    ("boxes_40hq", "boxes/40HQ", 12),
    ("cbm_40hq", "CBM/40HQ", 12),
    ("fob_40hq", "FOB $/40HQ", 14),
    ("ocean_per_box", "Ocean $/box ($10k)", 16),
    ("fob_x_duty", "FOB × 1.43", 12),
    ("landed_box", "Landed $/box", 14),
    ("landed_per_lf", "Landed $/LF", 12),
    ("us_sku", "US check SKU", 14),
    ("us_box", "US $/box", 12),
    ("us_per_lf", "US $/LF", 12),
    ("beat_vs_us", "Beat vs US $/box", 16),
    ("moq", "MOQ", 18),
    ("lead", "Lead time", 22),
    ("loading", "Loading", 28),
    ("load_note", "40HQ note", 40),
    ("verdict", "Verdict", 55),
]


def style_header(ws):
    for col, (_, title, width) in enumerate(COLS, 1):
        cell = ws.cell(1, col, title)
        cell.fill = HEADER
        cell.font = WHITE_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}1"


def write_rows(ws, rows):
    money = {
        "fob",
        "fob_per_lf",
        "fob_40hq",
        "ocean_per_box",
        "fob_x_duty",
        "landed_box",
        "landed_per_lf",
        "us_box",
        "us_per_lf",
        "beat_vs_us",
    }
    for r_i, row in enumerate(rows, 2):
        for c_i, (key, _, _) in enumerate(COLS, 1):
            val = row.get(key, "")
            cell = ws.cell(r_i, c_i, val)
            cell.alignment = WRAP
            cell.border = THIN
            cell.font = Font(name="Calibri", size=10)
            if key in money and isinstance(val, (int, float)):
                cell.number_format = "0.00" if abs(val) >= 1 else "0.00000"
            if key == "beat_vs_us" and isinstance(val, (int, float)) and val > 0:
                cell.fill = GREEN
            if key == "verdict":
                cell.fill = GREEN if "Beats" in str(val) or "Still beats" in str(val) else YELLOW
        ws.row_dimensions[r_i].height = 48


def notes_sheet(ws):
    lines = [
        "INTERNAL — do not forward to Fred / Halex.",
        "",
        "Source: cursor/product-pricing-analysis/halex/Carpet tackstrip Prices to All Pro Building Supplies LLC 20260909.xlsx",
        "Date on sheet: 9 Sep 2026. FOB Shanghai USD per box. Three SKUs. CARB/EPA Yes. Poplar 4 ft (1220 mm).",
        "100 pcs/box = 400 LF. MOQ = 40HQ mixed SKUs (not 100 boxes). Lead 2–3 weeks after deposit.",
        "Loading: palletized + hand stack to fill. SGR-120 1740 boxes / 73 CBM. SGR-340 and SGR-360 1540 boxes / 71.7 CBM.",
        "",
        "Landed = FOB × 1.43 (planning China duty stack) + $10,000 / boxes in that 40HQ.",
        "Confirm HTS with the broker (wood + nails). Do not invent a different duty. 1.43 is the house planning stack.",
        "",
        "US check is the public Home Depot Halex 100-pack (same mill brand, retail):",
        "  HD-320-P-8 wood 100-pack $28.65  https://www.homedepot.com/p/203301542",
        "  HD-340-P-8 concrete 100-pack $37.47  https://www.homedepot.com/p/203301576",
        "SGR-120 is 21.6 mm (~7/8\"); HD-320 is listed 1 in. Still the public wood box.",
        "",
        "Unlike Dongsheng carpet and Anerte pad, tackstrip still beats the US box at $10k ocean.",
        "No factory PO until the combined mill window (glass / carpet / pad / copper still open).",
        "Samples still outstanding — send Halex_Fred_quote_ack_SEND.txt. DHL to 1600 Livingston Ave, North Brunswick, NJ 08902.",
        "Anerte gripper stays backup only. Do not add storefront SKUs until the PO is locked.",
        "",
        "Rebuild: python3 cursor/product-pricing-analysis/halex/build_halex_compare.py",
    ]
    ws["A1"] = "Halex Fred quote — notes"
    ws["A1"].font = Font(bold=True, size=14, name="Calibri")
    for i, line in enumerate(lines, 3):
        ws.cell(i, 1, line).font = Font(name="Calibri", size=11)
        ws.row_dimensions[i].height = 18
    ws.column_dimensions["A"].width = 140


def main() -> None:
    rows = [landed(s) for s in SKUS]
    rows.append(mix_row())

    wb = Workbook()
    ws = wb.active
    ws.title = "Landed"
    style_header(ws)
    write_rows(ws, rows)
    notes = wb.create_sheet("Notes")
    notes_sheet(notes)
    wb.save(COMPARE_XLSX)

    with COMPARE_CSV.open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=[k for k, _, _ in COLS], extrasaction="ignore")
        w.writeheader()
        for row in rows:
            w.writerow({k: row.get(k, "") for k, _, _ in COLS})

    print(f"wrote {COMPARE_XLSX.name} and {COMPARE_CSV.name}")
    for row in rows:
        print(
            f"  {row['sku']}: FOB ${row['fob']}/box  landed ${row['landed_box']}/box  "
            f"${row['landed_per_lf']}/LF  vs US ${row['us_box']}  beat ${row['beat_vs_us']}  "
            f"40HQ FOB ${row['fob_40hq']}"
        )


if __name__ == "__main__":
    main()
