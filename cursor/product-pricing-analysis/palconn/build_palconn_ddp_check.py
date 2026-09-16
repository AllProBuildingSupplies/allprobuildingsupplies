#!/usr/bin/env python3
"""Palconn 15 Sep commercial invoice: DDP vs our FOB×1.43+$10k/40ft landed, vs site sell.

Internal — do not forward to Leela.
Invoice IN2026-0915 / Client PO PAL20260912. DDP labeled NJ 08902.
"""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
SITE = ROOT / "assets" / "products.csv"
OUT_XLSX = HERE / "Palconn_15Sep_DDP_vs_landed.xlsx"
OUT_CSV = HERE / "Palconn_15Sep_DDP_vs_landed.csv"
OUT_INV = HERE / "Palconn_15Sep_DDP_invoice.csv"

DUTY = 1.43  # same blunt stack as palconn/build_palconn_quote_compare.py
OCEAN = 10_000.0
CBM_40 = 67.7
OCEAN_PER_CBM = OCEAN / CBM_40

# Invoice IN2026-0915 line amounts (printed FOB Amount / DDP Amount). Unit = FOB Amount / qty.
# Pipe qty is feet. Fittings qty is pieces. CBM is the line total from the invoice.
LINES = [
    # family, code, desc, size, color, unit, qty, fob_amt, ddp_amt, pack, packages, meas, cbm
    ("PPSU", "PEX-CPLNG", "Coupling F2159 PPSU crimp", "1/2", "", "pc", 3000, 307, 479, 1000, 3, 0.012, 0.036),
    ("PPSU", "PEX-CPLNG", "Coupling F2159 PPSU crimp", "1", "", "pc", 3000, 1080, 1685, 200, 15, 0.012, 0.180),
    ("PPSU", "PEX-CPLNG", "Coupling F2159 PPSU crimp", "3/4", "", "pc", 3000, 582, 908, 500, 6, 0.012, 0.072),
    ("Copper ring", "PEX-F1807-RING", "Copper crimp ring", "1/2", "", "pc", 30000, 4155, 8061, 2000, 15, 0.012, 0.180),
    ("Copper ring", "PEX-F1807-RING", "Copper crimp ring", "1", "", "pc", 10000, 2447, 4747, 500, 20, 0.012, 0.240),
    ("Copper ring", "PEX-F1807-RING", "Copper crimp ring", "3/4", "", "pc", 20000, 3970, 7703, 1000, 20, 0.012, 0.240),
    ("PPSU", "PEX-ELBOW", "90 elbow F2159 PPSU crimp", "1/2", "", "pc", 9000, 1334, 2081, 600, 15, 0.012, 0.180),
    ("PPSU", "PEX-ELBOW", "90 elbow F2159 PPSU crimp", "1", "", "pc", 3000, 1779, 2775, 150, 20, 0.012, 0.240),
    ("PPSU", "PEX-ELBOW", "90 elbow F2159 PPSU crimp", "3/4", "", "pc", 4500, 1414, 2205, 300, 15, 0.012, 0.180),
    ("PPSU", "PEX-PLUG", "Plug F2159 PPSU", "1/2", "", "pc", 6000, 402, 628, 2000, 3, 0.012, 0.036),
    ("PPSU", "PEX-PLUG", "Plug F2159 PPSU", "1", "", "pc", 3000, 688, 1074, 600, 5, 0.012, 0.060),
    ("PPSU", "PEX-PLUG", "Plug F2159 PPSU", "3/4", "", "pc", 3000, 371, 578, 1000, 3, 0.012, 0.036),
    ("PPSU", "PEX-REDUCER", "Reducing coupling F2159 PPSU", "1 x 3/4", "", "pc", 3000, 868, 1354, 200, 15, 0.012, 0.180),
    ("PPSU", "PEX-REDUCER", "Reducing coupling F2159 PPSU", "3/4 x 1/2", "", "pc", 3000, 476, 743, 500, 6, 0.012, 0.072),
    ("PPSU", "PEX-REDTEE", "Reducing tee F2159 PPSU", "1/2 x 1/2 x 3/4", "", "pc", 3000, 826, 1288, 200, 15, 0.012, 0.180),
    ("PPSU", "PEX-REDTEE", "Reducing tee F2159 PPSU", "1 x 1 x 1/2", "", "pc", 3000, 1652, 2577, 150, 20, 0.012, 0.240),
    ("PPSU", "PEX-REDTEE", "Reducing tee F2159 PPSU", "1 x 1 x 3/4", "", "pc", 3000, 1906, 2973, 100, 30, 0.012, 0.360),
    ("PPSU", "PEX-REDTEE", "Reducing tee F2159 PPSU", "1 x 3/4 x 1", "", "pc", 3000, 2118, 3304, 150, 20, 0.012, 0.240),
    ("PPSU", "PEX-REDTEE", "Reducing tee F2159 PPSU", "1 x 3/4 x 3/4", "", "pc", 3000, 1694, 2643, 150, 20, 0.012, 0.240),
    ("PPSU", "PEX-REDTEE", "Reducing tee F2159 PPSU", "3/4 x 1/2 x 1/2", "", "pc", 3000, 900, 1404, 300, 10, 0.012, 0.120),
    ("PPSU", "PEX-REDTEE", "Reducing tee F2159 PPSU", "3/4 x 1/2 x 3/4", "", "pc", 3000, 1112, 1734, 250, 12, 0.012, 0.144),
    ("PPSU", "PEX-REDTEE", "Reducing tee F2159 PPSU", "3/4 x 3/4 x 1/2", "", "pc", 3000, 1059, 1652, 250, 12, 0.012, 0.144),
    ("PPSU", "PEX-REDTEE", "Reducing tee F2159 PPSU", "3/4 x 3/4 x 1", "", "pc", 3000, 1641, 2560, 100, 30, 0.012, 0.360),
    ("PPSU", "PEX-TEE", "Tee F2159 PPSU crimp", "1/2", "", "pc", 3000, 635, 991, 400, 7.5, 0.012, 0.090),
    ("PPSU", "PEX-TEE", "Tee F2159 PPSU crimp", "1", "", "pc", 3000, 2647, 4129, 100, 30, 0.012, 0.360),
    ("PPSU", "PEX-TEE", "Tee F2159 PPSU crimp", "3/4", "", "pc", 3000, 1271, 1982, 200, 15, 0.012, 0.180),
    ("PEX-B stick", "PEX-B-PIPE-ST-BLUE", "PEX-B potable pipe, 20 ft stick", "1", "Blue", "ft", 50000, 12800, 19968, "20ft/stick 20sticks/bag", 125, 0.09, 11.25),
    ("PEX-B stick", "PEX-B-PIPE-ST-RED", "PEX-B potable pipe, 20 ft stick", "1", "Red", "ft", 50000, 12800, 19968, "20ft/stick 20sticks/bag", 125, 0.09, 11.25),
    ("PEX-B stick", "PEX-B-PIPE-ST-RED", "PEX-B potable pipe, 20 ft stick", "3/4", "Red", "ft", 50000, 7850, 12246, "20ft/stick 25sticks/bag", 100, 0.09, 9.00),
    ("PEX-B stick", "PEX-B-PIPE-ST-BLUE", "PEX-B potable pipe, 20 ft stick", "3/4", "Blue", "ft", 50000, 7850, 12246, "20ft/stick 25sticks/bag", 100, 0.09, 9.00),
    ("PEX-B stick", "PEX-B-PIPE-ST-BLUE", "PEX-B potable pipe, 20 ft stick", "1/2", "Blue", "ft", 50000, 4050, 6318, "20ft/stick 50sticks/bag", 50, 0.09, 4.50),
    ("PEX-B stick", "PEX-B-PIPE-ST-RED", "PEX-B potable pipe, 20 ft stick", "1/2", "Red", "ft", 50000, 4050, 6318, "20ft/stick 50sticks/bag", 50, 0.09, 4.50),
    ("PEX-B coil", "PEX-B-PIPE-RL-BLUE", "PEX-B potable pipe, 100 ft coil", "1/2", "Blue", "ft", 50000, 4050, 6318, 100, 500, 0.025, 12.50),
    ("PEX-B coil", "PEX-B-PIPE-RL-RED", "PEX-B potable pipe, 100 ft coil", "1/2", "Red", "ft", 50000, 4050, 6318, 100, 500, 0.025, 12.50),
    ("PEX-B coil", "PEX-B-PIPE-RL-BLUE", "PEX-B potable pipe, 100 ft coil", "3/4", "Blue", "ft", 30000, 4710, 7348, 100, 300, 0.043, 12.90),
    ("PEX-B coil", "PEX-B-PIPE-RL-RED", "PEX-B potable pipe, 100 ft coil", "3/4", "Red", "ft", 30000, 4710, 7348, 100, 300, 0.043, 12.90),
    ("PEX-B coil", "PEX-B-PIPE-RL-BLUE", "PEX-B potable pipe, 100 ft coil", "1", "Blue", "ft", 30000, 7680, 11981, 100, 300, 0.064, 19.20),
    ("PEX-B coil", "PEX-B-PIPE-RL-RED", "PEX-B potable pipe, 100 ft coil", "1", "Red", "ft", 30000, 7680, 11981, 100, 300, 0.064, 19.20),
]

INV_FOB_PRINTED = 119_614
INV_DDP_PRINTED = 190_616
INV_CBM_PRINTED = 143.0
INV_PKG_PRINTED = 3133


def canon_size(raw) -> str:
    s = str(raw or "").replace('"', "").replace("″", "").replace("''", "")
    s = s.replace("×", "x").replace("X", "x").replace(" ", "")
    return s


def load_site():
    prices = {}
    with SITE.open(newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            code = (r.get("Code") or "").strip()
            size = canon_size(r.get("Size") or "")
            color = (r.get("Color") or "").strip()
            p = r.get("Price") or r.get("price") or ""
            try:
                price = float(str(p).replace("$", "").strip()) if str(p).strip() else None
            except ValueError:
                price = None
            if code and price is not None:
                prices[(code, size, color)] = price
                prices.setdefault((code, size, ""), price)
    return prices


def money(v, nd=4):
    if v is None:
        return None
    return round(float(v), nd)


def gm_pct(sell, cost):
    if sell is None or cost is None or sell <= 0:
        return None
    return round(100.0 * (sell - cost) / sell, 1)


def band(gm):
    if gm is None:
        return "no site price"
    if gm < 0:
        return "UNDERWATER — raise sell or skip"
    if gm < 40:
        return "thin — below 50% range"
    if gm < 45:
        return "shy of 50% range"
    if gm <= 55:
        return "in 50% range"
    if gm <= 65:
        return "fat — above 50% range"
    return "very fat — room to cut"


def build_rows(site):
    rows = []
    for fam, code, desc, size, color, unit, qty, fob_amt, ddp_amt, pack, pkgs, meas, cbm in LINES:
        fob_u = fob_amt / qty
        ddp_u = ddp_amt / qty
        ocean_u = (cbm / qty) * OCEAN_PER_CBM if qty else 0.0
        our_u = fob_u * DUTY + ocean_u
        our_line = fob_amt * DUTY + cbm * OCEAN_PER_CBM
        ddp_vs_our = 100.0 * (ddp_u - our_u) / our_u if our_u else None
        ddp_over_fob = ddp_u / fob_u if fob_u else None

        site_key = (code, canon_size(size), color)
        site_raw = site.get(site_key) or site.get((code, canon_size(size), ""))

        # Site pipe SKUs are priced per stick (20 ft) or per coil (100 ft), Palconn is $/ft.
        if unit == "ft" and "stick" in desc:
            site_per_invoice_unit = (site_raw / 20.0) if site_raw is not None else None
            site_sell_unit = "20 ft stick"
            site_sell_price = site_raw
            ddp_sell_unit = ddp_u * 20.0
            our_sell_unit = our_u * 20.0
            fob_sell_unit = fob_u * 20.0
        elif unit == "ft" and "coil" in desc:
            site_per_invoice_unit = (site_raw / 100.0) if site_raw is not None else None
            site_sell_unit = "100 ft coil"
            site_sell_price = site_raw
            ddp_sell_unit = ddp_u * 100.0
            our_sell_unit = our_u * 100.0
            fob_sell_unit = fob_u * 100.0
        else:
            site_per_invoice_unit = site_raw
            site_sell_unit = "pc"
            site_sell_price = site_raw
            ddp_sell_unit = ddp_u
            our_sell_unit = our_u
            fob_sell_unit = fob_u

        gm_ddp = gm_pct(site_per_invoice_unit, ddp_u)
        gm_our = gm_pct(site_per_invoice_unit, our_u)
        target_50_ddp = 2.0 * ddp_sell_unit  # 50% GM on Palconn DDP, in the site selling unit
        target_50_our = 2.0 * our_sell_unit

        rows.append(
            {
                "family": fam,
                "code": code,
                "desc": desc,
                "size": size,
                "color": color,
                "unit": unit,
                "qty": qty,
                "fob_unit": money(fob_u, 4),
                "ddp_unit": money(ddp_u, 4),
                "our_landed_unit": money(our_u, 4),
                "ocean_unit": money(ocean_u, 4),
                "ddp_over_fob": money(ddp_over_fob, 3),
                "ddp_vs_our_pct": money(ddp_vs_our, 1),
                "fob_amt": fob_amt,
                "ddp_amt": ddp_amt,
                "our_landed_amt": money(our_line, 2),
                "cbm": cbm,
                "packages": pkgs,
                "pack": pack,
                "site_sell_unit": site_sell_unit,
                "site_sell": site_sell_price,
                "fob_in_sell_unit": money(fob_sell_unit, 4),
                "ddp_in_sell_unit": money(ddp_sell_unit, 4),
                "our_in_sell_unit": money(our_sell_unit, 4),
                "gm_on_ddp_pct": gm_ddp,
                "gm_on_our_pct": gm_our,
                "target_sell_50gm_ddp": money(target_50_ddp, 2),
                "target_sell_50gm_our": money(target_50_our, 2),
                "band_ddp": band(gm_ddp),
            }
        )
    return rows


HEADER_FILL = PatternFill("solid", fgColor="C00000")
WHITE = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
GREEN = PatternFill("solid", fgColor="C6EFCE")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
RED = PatternFill("solid", fgColor="F8CBAD")
BLUE = PatternFill("solid", fgColor="DDEBF7")
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
WRAP = Alignment(wrap_text=True, vertical="center")

COLS = [
    ("family", "Family", 14),
    ("code", "APBS code", 22),
    ("desc", "Description", 34),
    ("size", "Size", 16),
    ("color", "Color", 8),
    ("unit", "Inv unit", 8),
    ("qty", "Qty", 10),
    ("fob_unit", "Palconn FOB $/unit", 14),
    ("ddp_unit", "Palconn DDP $/unit", 14),
    ("our_landed_unit", "Our landed $/unit", 14),
    ("ocean_unit", "Our ocean $/unit", 12),
    ("ddp_over_fob", "DDP / FOB", 10),
    ("ddp_vs_our_pct", "DDP vs our %", 12),
    ("fob_amt", "FOB $", 12),
    ("ddp_amt", "DDP $", 12),
    ("our_landed_amt", "Our landed $", 12),
    ("cbm", "CBM", 8),
    ("site_sell_unit", "Site unit", 14),
    ("site_sell", "Site sell $", 12),
    ("ddp_in_sell_unit", "DDP in site unit", 14),
    ("our_in_sell_unit", "Our landed in site unit", 16),
    ("gm_on_ddp_pct", "GM% on Palconn DDP", 14),
    ("gm_on_our_pct", "GM% on our landed", 14),
    ("target_sell_50gm_ddp", "Sell for 50% GM (DDP)", 16),
    ("band_ddp", "50% range vs DDP", 28),
]


def write_xlsx(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Line check"
    for i, (_, title, width) in enumerate(COLS, 1):
        c = ws.cell(1, i, title)
        c.fill = HEADER_FILL
        c.font = WHITE
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{1 + len(rows)}"

    money_cols = {
        "fob_unit",
        "ddp_unit",
        "our_landed_unit",
        "ocean_unit",
        "fob_amt",
        "ddp_amt",
        "our_landed_amt",
        "site_sell",
        "ddp_in_sell_unit",
        "our_in_sell_unit",
        "target_sell_50gm_ddp",
        "fob_in_sell_unit",
    }
    for ri, row in enumerate(rows, 2):
        for ci, (k, _, _) in enumerate(COLS, 1):
            v = row.get(k)
            cell = ws.cell(ri, ci, v)
            cell.alignment = WRAP
            cell.border = THIN
            cell.font = Font(name="Calibri", size=9)
            if k in money_cols and isinstance(v, (int, float)):
                cell.number_format = '"$"#,##0.00' if abs(v) >= 1 else '"$"0.000'
            if k == "ddp_vs_our_pct" and isinstance(v, (int, float)):
                cell.number_format = '0.0"%"'
                if abs(v) <= 5:
                    cell.fill = GREEN
                elif abs(v) <= 15:
                    cell.fill = YELLOW
                else:
                    cell.fill = RED
            if k == "band_ddp":
                t = str(v)
                if "in 50%" in t:
                    cell.fill = GREEN
                elif "UNDERWATER" in t or "thin" in t:
                    cell.fill = RED
                elif "shy" in t or "fat" in t:
                    cell.fill = YELLOW
            if k == "gm_on_ddp_pct" and isinstance(v, (int, float)):
                cell.number_format = '0.0"%"'
                if v < 0:
                    cell.fill = RED
                elif 45 <= v <= 55:
                    cell.fill = GREEN
                elif v < 45:
                    cell.fill = YELLOW
                else:
                    cell.fill = BLUE
        ws.row_dimensions[ri].height = 28

    tot = 2 + len(rows)
    ws.cell(tot, 1, "INVOICE TOTAL").font = Font(bold=True)
    ws.cell(tot, 7, sum(r["qty"] for r in rows))
    ws.cell(tot, 14, sum(r["fob_amt"] for r in rows)).number_format = '"$"#,##0.00'
    ws.cell(tot, 15, sum(r["ddp_amt"] for r in rows)).number_format = '"$"#,##0.00'
    ws.cell(tot, 16, sum(r["our_landed_amt"] for r in rows)).number_format = '"$"#,##0.00'
    ws.cell(tot, 17, sum(r["cbm"] for r in rows))
    for c in range(1, len(COLS) + 1):
        ws.cell(tot, c).fill = PatternFill("solid", fgColor="D6E3F0")
        ws.cell(tot, c).font = Font(bold=True, name="Calibri", size=9)
        ws.cell(tot, c).border = THIN

    notes = wb.create_sheet("Notes")
    fob_sum = sum(r["fob_amt"] for r in rows)
    ddp_sum = sum(r["ddp_amt"] for r in rows)
    our_sum = sum(r["our_landed_amt"] for r in rows)
    cbm_sum = sum(r["cbm"] for r in rows)
    lines = [
        "INTERNAL — do not forward to Palconn / Leela.",
        "",
        "Source: Palconn commercial invoice IN2026-0915, 15 Sep 2026, client PO PAL20260912.",
        "FOB Qingdao Port vs DDP 1600 Livingston Ave, North Brunswick NJ 08902.",
        "Payment 30% TT before production / 70% before shipment. NSF/UPC/CSA.",
        "",
        "Our estimated landed (same method as Palconn_Leela_quote_COMPARE):",
        "  Landed = FOB × 1.43 + ($10,000 / 67.7 CBM) × line CBM.",
        "  1.43 = MFN + China Section 301 ~25% + Jul 2026 overlay ~12.5% (blunt plastic stack).",
        "  Ocean is allocated by cube, not by value.",
        "",
        f"Printed invoice totals: FOB ${INV_FOB_PRINTED:,}  DDP ${INV_DDP_PRINTED:,}  CBM {INV_CBM_PRINTED}  pkgs {INV_PKG_PRINTED}.",
        f"Transcribed sums:       FOB ${fob_sum:,.0f}  DDP ${ddp_sum:,.0f}  CBM {cbm_sum:.2f}.",
        f"Our landed on this mix: ${our_sum:,.2f}  ({100 * ddp_sum / our_sum:.1f}% of Palconn DDP).",
        "",
        "Palconn DDP / FOB is ~1.56× on plastics (pipe + PPSU) and ~1.94× on copper rings.",
        "They did not allocate ocean by CBM: 20 ft stick and 100 ft coil have the same DDP/ft.",
        "Copper rings: Palconn DDP implies ~93.5% all-in. That fits MFN + 301 + overlay + Section 232 copper ~50%,",
        "which our 1.43 plastic stack does not include. Treat Palconn DDP as the better ring cost until a broker confirms HTS.",
        "",
        "True DDP to the warehouse: Baruch confirmed 16 Sep her DDP includes delivery to Livingston — no dest add-on.",
        "Palconn DDP sitting ~1% under FOB×1.43+ocean means the formula already approximates all-in warehouse cost.",
        "Sell cards = Gator/NJPD (quote 11237145, ack 11234952), rounded up on the site. Rings lose money at Gator; sticks/PPSU ~25–38% GM.",
        "",
        "Site sell = assets/products.csv. Pipe SKUs are priced per 20 ft stick or per 100 ft coil, not per foot.",
        "Gross margin % = (site − cost) / site. 50% range = 45–55% GM (sell ≈ 2× cost).",
        "",
        "Rebuild: python3 cursor/product-pricing-analysis/palconn/build_palconn_ddp_check.py",
    ]
    notes["A1"] = "Palconn 15 Sep DDP vs our landed"
    notes["A1"].font = Font(bold=True, size=14, name="Calibri")
    for i, line in enumerate(lines, 3):
        notes.cell(i, 1, line).font = Font(name="Calibri", size=11)
    notes.column_dimensions["A"].width = 140
    wb.save(OUT_XLSX)


def main():
    site = load_site()
    rows = build_rows(site)

    fob_sum = sum(r["fob_amt"] for r in rows)
    ddp_sum = sum(r["ddp_amt"] for r in rows)
    our_sum = sum(r["our_landed_amt"] for r in rows)
    cbm_sum = sum(r["cbm"] for r in rows)
    pkg_sum = sum(r["packages"] for r in rows)

    inv_fields = [
        "family",
        "code",
        "desc",
        "size",
        "color",
        "unit",
        "qty",
        "fob_amt",
        "ddp_amt",
        "pack",
        "packages",
        "cbm",
    ]
    with OUT_INV.open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=inv_fields, extrasaction="ignore")
        w.writeheader()
        for fam, code, desc, size, color, unit, qty, fob_amt, ddp_amt, pack, pkgs, meas, cbm in LINES:
            w.writerow(
                {
                    "family": fam,
                    "code": code,
                    "desc": desc,
                    "size": size,
                    "color": color,
                    "unit": unit,
                    "qty": qty,
                    "fob_amt": fob_amt,
                    "ddp_amt": ddp_amt,
                    "pack": pack,
                    "packages": pkgs,
                    "cbm": cbm,
                }
            )

    with OUT_CSV.open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=[k for k, _, _ in COLS], extrasaction="ignore")
        w.writeheader()
        for row in rows:
            w.writerow({k: row.get(k, "") for k, _, _ in COLS})

    write_xlsx(rows)

    print(f"wrote {OUT_XLSX.name}")
    print(f"transcribed FOB ${fob_sum:,.0f} vs printed ${INV_FOB_PRINTED:,}  delta {fob_sum - INV_FOB_PRINTED:.0f}")
    print(f"transcribed DDP ${ddp_sum:,.0f} vs printed ${INV_DDP_PRINTED:,}  delta {ddp_sum - INV_DDP_PRINTED:.0f}")
    print(f"CBM {cbm_sum:.2f} vs printed {INV_CBM_PRINTED}  pkgs {pkg_sum} vs {INV_PKG_PRINTED}")
    print(f"our landed ${our_sum:,.2f}  Palconn DDP is {100 * ddp_sum / our_sum:.1f}% of our number")
    print(f"DDP/FOB {ddp_sum / fob_sum:.3f}")

    print("\nFamily DDP vs our landed:")
    fams = {}
    for r in rows:
        d = fams.setdefault(r["family"], {"ddp": 0, "our": 0, "fob": 0, "n": 0})
        d["ddp"] += r["ddp_amt"]
        d["our"] += r["our_landed_amt"]
        d["fob"] += r["fob_amt"]
        d["n"] += 1
    for fam, d in fams.items():
        vs = 100 * (d["ddp"] - d["our"]) / d["our"]
        print(f"  {fam}: FOB ${d['fob']:,.0f}  DDP ${d['ddp']:,.0f}  our ${d['our']:,.0f}  DDP vs our {vs:+.1f}%")

    print("\n50% GM vs Palconn DDP (unique size/code, colors collapsed):")
    seen = set()
    counts = {"in 50% range": 0, "thin — below 50% range": 0, "shy of 50% range": 0, "fat — above 50% range": 0, "very fat — room to cut": 0, "UNDERWATER — raise sell or skip": 0}
    for r in rows:
        k = (r["code"], r["size"])
        if k in seen:
            continue
        seen.add(k)
        print(
            f"  {r['code']} {r['size']}: site {r['site_sell']} / {r['site_sell_unit']}  "
            f"DDP {r['ddp_in_sell_unit']}  GM {r['gm_on_ddp_pct']}%  "
            f"need {r['target_sell_50gm_ddp']}  | {r['band_ddp']}"
        )
        counts[r["band_ddp"]] = counts.get(r["band_ddp"], 0) + 1
    print("\nSKU bands:", {k: v for k, v in counts.items() if v})


if __name__ == "__main__":
    main()
