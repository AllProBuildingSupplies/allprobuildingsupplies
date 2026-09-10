#!/usr/bin/env python3
"""Land Palconn/Leela 10 Sep quote vs Tommur / Zhenpeng. Internal — do not send."""

from __future__ import annotations

import csv
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
QUOTE_10 = HERE / "20260910 Palconn_Full_Catalog_RFQ_SEND.xlsx"
QUOTE_8 = HERE / "20260908 Palconn_Full_Catalog_RFQ_SEND.xlsx"
FACTORY = ROOT / "cursor" / "product-pricing-analysis" / "Factory_Order_PVC_PEX_45HQ.csv"
SITE = ROOT / "assets" / "products.csv"
OUT_XLSX = HERE / "Palconn_Leela_quote_COMPARE.xlsx"
OUT_CSV = HERE / "Palconn_Leela_quote_COMPARE.csv"

DUTY = 1.43
OCEAN = 10_000.0
CBM_40 = 67.7  # same as build_tommur_cost_margin.py
OCEAN_PER_CBM = OCEAN / CBM_40

ZHENPENG = {
    ("PEX-ELBOW", "1/2"): 0.167,
    ("PEX-ELBOW", "3/4"): 0.342,
    ("PEX-ELBOW", "1"): 0.574,
    ("PEX-CPLNG", "1/2"): 0.104,
    ("PEX-CPLNG", "3/4"): 0.201,
    ("PEX-CPLNG", "1"): 0.427,
    ("PEX-TEE", "1/2"): 0.230,
    ("PEX-TEE", "3/4"): 0.474,
    ("PEX-TEE", "1"): 0.910,
    ("PEX-REDTEE", "3/4x3/4x1/2"): 0.441,
}

HEADER = PatternFill("solid", fgColor="C00000")
WHITE = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
GREEN = PatternFill("solid", fgColor="C6EFCE")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
RED = PatternFill("solid", fgColor="F8CBAD")
WRAP = Alignment(wrap_text=True, vertical="center")
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)

COLS = [
    ("family", "Family", 20),
    ("code", "Code", 18),
    ("desc", "Description", 36),
    ("size", "Size", 16),
    ("unit", "Unit", 12),
    ("fob8", "FOB 8 Sep", 12),
    ("fob10", "FOB 10 Sep", 12),
    ("fob_delta_pct", "FOB Δ%", 10),
    ("moq", "MOQ", 10),
    ("pack", "Pack", 22),
    ("cbm", "CBM/carton", 12),
    ("gw", "GW kg/carton", 12),
    ("lead", "Lead", 8),
    ("nsf", "NSF", 12),
    ("units_carton", "Units/carton", 12),
    ("ocean_unit", "Ocean $/unit", 12),
    ("landed", "Landed $ (1.43+$10k)", 16),
    ("our_fob", "Our FOB", 12),
    ("our_src", "Our FOB source", 18),
    ("tommur_ddp", "Tommur DDP", 12),
    ("zhenpeng", "Zhenpeng REV", 12),
    ("site", "Site sell", 10),
    ("verdict", "Verdict", 42),
]


def canon_size(raw) -> str:
    s = str(raw or "").replace('"', "").replace("″", "").replace("''", "")
    s = s.replace("×", "x").replace("X", "x").replace(" ", "")
    return s


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    if s in ("//", "-", "n/a", "N/A", "None"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def load_quote(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb["Quote"] if "Quote" in wb.sheetnames else wb.active
    start = 2
    rows = []
    for r in range(start, ws.max_row + 1):
        family = ws.cell(r, 1).value
        if not family:
            continue
        rows.append(
            {
                "family": family,
                "code": ws.cell(r, 2).value,
                "desc": ws.cell(r, 3).value,
                "size": ws.cell(r, 4).value,
                "unit": ws.cell(r, 6).value,
                "fob": num(ws.cell(r, 8).value),
                "moq": ws.cell(r, 9).value,
                "pack": ws.cell(r, 10).value,
                "cbm": num(ws.cell(r, 11).value),
                "gw": num(ws.cell(r, 12).value),
                "lead": ws.cell(r, 13).value,
                "nsf": ws.cell(r, 14).value,
            }
        )
    return rows


def key(r):
    return (str(r["family"]), str(r["code"]), canon_size(r["size"]), str(r["desc"]))


def units_per_carton(r) -> float | None:
    pack = r.get("pack")
    unit = str(r.get("unit") or "")
    if pack is None or pack == "" or str(pack).strip() == "//":
        return None
    if isinstance(pack, str):
        m = re.search(r"(\d+)\s*sticks", pack, re.I)
        if m and "ft" in unit:
            return int(m.group(1)) * 20.0
        n = num(pack)
        if n is None:
            return None
        pack = n
    pack = float(pack)
    if unit == "20 ft stick" and abs(pack - 20) < 1e-6:
        # Palconn pack=20 and GW match one 20 ft length.
        return 1.0
    return pack


def ocean_per_unit(r) -> float | None:
    cbm = r.get("cbm")
    upc = units_per_carton(r)
    if not cbm or not upc:
        return None
    return OCEAN_PER_CBM * cbm / upc


def load_factory():
    fob, ddp = {}, {}
    with FACTORY.open(newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            code = (r.get("APBS_Code") or "").strip()
            size = canon_size(r.get("Size") or "")
            k = (code, size)
            fn, dn = num(r.get("FOB_USD")), num(r.get("DDP_USD"))
            if fn is not None:
                fob[k] = fn
            if dn is not None:
                ddp[k] = dn
    return fob, ddp


def load_site():
    prices = {}
    with SITE.open(newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            code = (r.get("Code") or "").strip()
            size = canon_size(r.get("Size") or "")
            p = num(r.get("Price") or r.get("price"))
            if code and p is not None:
                prices[(code, size)] = p
    return prices


def verdict(row, fob10, our_fob, zhen, ocean, landed) -> str:
    if fob10 is None:
        if row["family"] in ("PVC DWV fittings", "CPVC SDR-11 (if manufactured)", "CPVC D2846 fittings (if manufactured)"):
            return "Omitted from 10 Sep sheet — still not quoted"
        return "Not quoted"
    if row["family"].startswith("PVC Sch") or row["family"].startswith("PVC foam"):
        site = row.get("site")
        if landed and site and row["unit"] == "20 ft stick" and site >= 5:
            if landed > site:
                return "Landed above site sell — keep Tommur"
            return "Landed under site (tight) — still Tommur unless FOB pipe exists"
        return "CBM in; keep Tommur for PVC pipe"
    if zhen is not None:
        if fob10 < zhen - 1e-6:
            return "Palconn cheaper FOB than Zhenpeng"
        if fob10 > zhen + 1e-6:
            return "Zhenpeng cheaper FOB"
        return "Match Zhenpeng"
    if our_fob is not None and row["our_src"] == "Tommur FOB":
        if fob10 < our_fob - 1e-6:
            return "Palconn cheaper FOB than Tommur"
        return "Tommur cheaper FOB"
    if row["family"].startswith("PEX F1807") or row["family"].startswith("PEX outlet"):
        return "Palconn only mill"
    if row["family"] == "PEX-B pipe":
        return "Palconn cheaper FOB than Tommur"
    if our_fob is None:
        return "No mill compare (new line)"
    return "vs Tommur DDP (not FOB)"


def roundish(v, nd=4):
    if v is None:
        return None
    return round(float(v), nd)


def build_rows():
    r10 = load_quote(QUOTE_10)
    r8 = load_quote(QUOTE_8)
    m8 = {key(r): r for r in r8}
    m10 = {key(r): r for r in r10}
    fob_map, ddp_map = load_factory()
    site_map = load_site()

    keys = list(dict.fromkeys(list(m10.keys()) + list(m8.keys())))
    out = []
    for k in keys:
        a = m10.get(k) or {}
        b = m8.get(k) or {}
        row = {
            "family": a.get("family") or b.get("family"),
            "code": a.get("code") or b.get("code"),
            "desc": a.get("desc") or b.get("desc"),
            "size": a.get("size") or b.get("size"),
            "unit": a.get("unit") or b.get("unit"),
            "fob8": b.get("fob"),
            "fob10": a.get("fob") if a else None,
            "moq": a.get("moq") if a else b.get("moq"),
            "pack": a.get("pack") if a else b.get("pack"),
            "cbm": a.get("cbm") if a else b.get("cbm"),
            "gw": a.get("gw") if a else b.get("gw"),
            "lead": a.get("lead") if a else b.get("lead"),
            "nsf": a.get("nsf") if a else b.get("nsf"),
        }
        f10, f8 = row["fob10"], row["fob8"]
        row["fob_delta_pct"] = round(100.0 * (f10 - f8) / f8, 1) if f10 and f8 else None
        row["units_carton"] = units_per_carton(row)
        row["ocean_unit"] = roundish(ocean_per_unit(row), 4)
        if f10 is not None:
            duty = f10 * DUTY
            ocean = row["ocean_unit"] or 0.0
            row["landed"] = roundish(duty + ocean, 4)
        else:
            row["landed"] = None
        ck = (row["code"], canon_size(row["size"]))
        row["our_fob"] = fob_map.get(ck)
        row["tommur_ddp"] = ddp_map.get(ck)
        row["zhenpeng"] = ZHENPENG.get(ck)
        row["site"] = site_map.get(ck)
        if row["zhenpeng"] is not None:
            row["our_src"] = "Zhenpeng REV FOB"
            if row["our_fob"] is None:
                row["our_fob"] = row["zhenpeng"]
        elif row["our_fob"] is not None:
            row["our_src"] = "Tommur FOB"
        else:
            row["our_src"] = ""
        row["verdict"] = verdict(row, f10, row["our_fob"], row["zhenpeng"], row["ocean_unit"], row["landed"])
        out.append(row)
    return out


def write_xlsx(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Landed"
    for i, (_, title, width) in enumerate(COLS, 1):
        c = ws.cell(1, i, title)
        c.fill = HEADER
        c.font = WHITE
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}1"
    money = {
        "fob8",
        "fob10",
        "ocean_unit",
        "landed",
        "our_fob",
        "tommur_ddp",
        "zhenpeng",
        "site",
        "cbm",
        "gw",
        "fob_delta_pct",
    }
    for ri, row in enumerate(rows, 2):
        for ci, (k, _, _) in enumerate(COLS, 1):
            v = row.get(k)
            cell = ws.cell(ri, ci, v)
            cell.alignment = WRAP
            cell.border = THIN
            cell.font = Font(name="Calibri", size=9)
            if k in money and isinstance(v, (int, float)):
                cell.number_format = "0.00" if abs(v) >= 1 else "0.0000"
            if k == "verdict":
                t = str(v)
                if "Palconn cheaper" in t or "Palconn only" in t:
                    cell.fill = GREEN
                elif "Zhenpeng cheaper" in t or "Tommur cheaper" in t or "above site" in t:
                    cell.fill = RED
                elif "not quoted" in t.lower() or "Omitted" in t:
                    cell.fill = YELLOW
        ws.row_dimensions[ri].height = 32

    notes = wb.create_sheet("Notes")
    lines = [
        "INTERNAL — do not forward to Palconn / Leela.",
        "",
        "Source 10 Sep: palconn/20260910 Palconn_Full_Catalog_RFQ_SEND.xlsx (84 lines: PEX + PVC pipe only).",
        "Source 8 Sep: palconn/20260908 Palconn_Full_Catalog_RFQ_SEND.xlsx (167 lines including DWV blanks).",
        "Ocean: ($10,000 / 67.7 CBM) × CBM_per_carton / units_per_carton. Same 40' cube as Tommur tracker.",
        "Duty stack 1.43. Landed = FOB × 1.43 + ocean/unit.",
        "PVC pipe: pack=20 and GW match one 20 ft stick, so units/carton = 1.",
        "PEX sticks: 50/25/20 sticks per bag × 20 ft.",
        "10 Sep changes: PEX-B −1.2%; all F2159 PPSU −5%; six extra reducing tees; PVC pipe MOQ/CBM/GW filled.",
        "F1807 unchanged. 1\" F2159 female adapter still blank. DWV + CPVC omitted from this return.",
        "No factory PO until the combined mill lock.",
        "",
        "Rebuild: python3 cursor/product-pricing-analysis/palconn/build_palconn_quote_compare.py",
    ]
    notes["A1"] = "Palconn 10 Sep quote — notes"
    notes["A1"].font = Font(bold=True, size=14, name="Calibri")
    for i, line in enumerate(lines, 3):
        notes.cell(i, 1, line).font = Font(name="Calibri", size=11)
    notes.column_dimensions["A"].width = 140
    wb.save(OUT_XLSX)


def main():
    rows = build_rows()
    write_xlsx(rows)
    with OUT_CSV.open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=[k for k, _, _ in COLS], extrasaction="ignore")
        w.writeheader()
        for row in rows:
            w.writerow({k: row.get(k, "") for k, _, _ in COLS})

    priced = [r for r in rows if r["fob10"] is not None]
    print(f"wrote {OUT_XLSX.name} ({len(rows)} rows, {len(priced)} priced on 10 Sep)")
    for fam in [
        "PEX-B pipe",
        "PEX F2159 PPSU",
        "PEX F1807 brass",
        "PEX outlet boxes",
        "PVC Sch40 pipe",
        "PVC foam-core pipe",
        "PVC DWV fittings",
    ]:
        sub = [r for r in rows if r["family"] == fam]
        n10 = sum(1 for r in sub if r["fob10"] is not None)
        print(f"  {fam}: {n10}/{len(sub)} priced")

    print("\nPEX-B stick landed:")
    for r in rows:
        if r["family"] == "PEX-B pipe" and "20 ft" in str(r["desc"]):
            print(f"  {r['size']} FOB {r['fob10']} landed {r['landed']} vs Tommur {r['our_fob']}")
    print("\nPPSU elbows:")
    for r in rows:
        if r["code"] == "PEX-ELBOW":
            print(f"  {r['size']} Pal {r['fob10']} Zhen {r['zhenpeng']} landed {r['landed']} | {r['verdict']}")
    print("\nPVC pipe landed vs site:")
    for r in rows:
        if r["unit"] == "20 ft stick" and r["fob10"]:
            print(
                f"  {r['code']} {r['size']} FOB {r['fob10']:.4f} CBM {r['cbm']} ocean {r['ocean_unit']} "
                f"landed {r['landed']} site {r['site']} | {r['verdict']}"
            )


if __name__ == "__main__":
    main()
