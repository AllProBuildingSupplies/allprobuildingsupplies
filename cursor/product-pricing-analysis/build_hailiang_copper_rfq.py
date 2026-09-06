#!/usr/bin/env python3
"""Build Hailiang/Ailiang copper RFQ from APBS products.csv + factory catalog pages 19-25."""

from __future__ import annotations

import csv
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

ROOT = Path(__file__).resolve().parent
CROPS = ROOT / "hailiang_catalog_crops"
SITE_CSV = Path("/workspace/assets/products.csv")
SITE_IMG = Path("/workspace/images")
OUT_XLSX = ROOT / "Hailiang_Copper_Brass_RFQ.xlsx"
OUT_CSV = ROOT / "Hailiang_Copper_Brass_RFQ.csv"

# Nominal inch (CTS) → Hailiang catalog φ mm (EN 1057). 3" and 4" are metric, not ASTM B88 OD.
INCH_TO_MM = {
    "1/2": 15,
    "3/4": 22,
    "1": 28,
    "1-1/4": 35,
    "1-1/2": 42,
    "2": 54,
    "2-1/2": 67,
    "3": 76,
    "4": 108,
}
MM_TO_INCH = {v: k for k, v in INCH_TO_MM.items()}
ASTM_OD_MM = {
    "1/2": 15.875,
    "3/4": 22.225,
    "1": 28.575,
    "1-1/4": 34.925,
    "1-1/2": 41.275,
    "2": 53.975,
    "2-1/2": 66.675,
    "3": 79.375,
    "4": 104.775,
}

S101 = {15, 22, 28, 35, 42, 54, 67, 76, 108, 133, 159, 219, 273, 325}
S102 = {
    (22, 15), (28, 15), (28, 22), (35, 15), (35, 22), (35, 28),
    (42, 22), (42, 28), (42, 35), (54, 28), (54, 35), (54, 42),
    (67, 35), (67, 42), (67, 54), (76, 42), (76, 54), (76, 67),
    (108, 54), (108, 67), (108, 76), (133, 67), (133, 76), (133, 108),
    (159, 76), (159, 108), (159, 133), (219, 108), (219, 133), (219, 159),
}
L101 = set(S101)
T101 = set(S101)
T102 = {
    (22, 15, 22), (28, 15, 28), (28, 22, 28), (35, 22, 35), (35, 28, 35),
    (42, 22, 42), (42, 28, 42), (42, 35, 42), (54, 28, 54), (54, 35, 54), (54, 42, 54),
    (67, 35, 67), (67, 42, 67), (67, 54, 67), (76, 42, 76), (76, 54, 76), (76, 67, 76),
    (108, 54, 108), (108, 67, 108), (108, 76, 108), (133, 67, 133), (133, 76, 133),
    (133, 108, 133), (159, 76, 159), (159, 108, 159), (159, 133, 159), (219, 159, 219),
}
T103 = {(15, 22, 15), (22, 28, 22), (28, 35, 28)}
T104 = {
    (22, 15, 15), (28, 22, 22), (28, 22, 15), (28, 15, 22),
    (28, 28, 22), (28, 28, 15), (35, 35, 22),
}

FAMILY = {
    "S-101": ("Straight Coupling", "直接头", 19, "S-101.png"),
    "S-102": ("Reducing Coupling", "异径直接头", 19, "S-102.png"),
    "S-103": ("Fitting Reduced (spigot)", "插口异径直接头", 19, "S-103.png"),
    "L-101": ("90° Elbow", "90°弯头", 19, "L-101.png"),
    "L-102": ("Reducing 90° Elbow", "异径90°弯头", 20, "L-102.png"),
    "L-103": ("45° Elbow", "45°弯头", 20, "L-103.png"),
    "T-101": ("Equal Tee", "正三通", 20, "T-101.png"),
    "T-102": ("Tee — reduced branch", "异径(中小)三通", 21, "T-102.png"),
    "T-103": ("Tee — reduced both ends (bullhead)", "异径(中大)三通", 20, "T-103.png"),
    "T-104": ("Tee — reduced end and branch", "异径(侧)三通", 20, "T-104.png"),
}


def parse_inch_parts(size: str) -> list[str]:
    s = (size or "").replace('"', "").strip()
    if not s:
        return []
    return [p.strip() for p in re.split(r"\s*x\s*", s, flags=re.I) if p.strip()]


def to_mm(parts: list[str]) -> list[int] | None:
    out = []
    for p in parts:
        if p not in INCH_TO_MM:
            return None
        out.append(INCH_TO_MM[p])
    return out


def fmt_phi(mms: list[int]) -> str:
    if len(mms) == 1:
        return f"φ{mms[0]}"
    return "φ" + "×".join(str(x) for x in mms)


def astm_note(parts: list[str]) -> str:
    flags = []
    for p in parts:
        if p in ("3", "4"):
            flags.append(
                f'{p}" CTS OD is {ASTM_OD_MM[p]:.3f} mm (ASTM B88); catalog lists φ{INCH_TO_MM[p]} (EN 1057). Confirm they make US tube size.'
            )
    return " ".join(flags)


def match_fitting(code: str, desc: str, parts: list[str], mms: list[int] | None) -> dict:
    d = desc.upper()
    c = code.upper()
    if mms is None:
        return unmatched("Size did not map to catalog φ mm.")

    ask3 = astm_note(parts)

    if "PIPE" in d or c.endswith("-PIPE"):
        return unmatched(
            "This PDF is fittings only (catalog pp. 19–25). Hailiang makes copper tube — request ASTM B88 Type L hard 10 ft sticks and Type K (soft coil / hard) separately."
        )
    if "PEX" in d or "PEXADPTR" in c:
        return unmatched("Copper-to-PEX adapter is not on catalog pages 19–25. Ask if they mold/machine this (F1807 crimp or press).")
    if "STUB" in d:
        return unmatched("Stub-out elbow is not on catalog pages 19–25. Ask if they make 1/2\" Type L 7×6 (or similar) stub-outs.")

    if "ELBOW 90" in d or c.endswith("-90"):
        if len(mms) != 1:
            return unmatched("90° elbow expected one size.")
        if mms[0] in L101:
            return hit("L-101", [mms[0]], ask3)
        return unmatched(f"φ{mms[0]} not listed on L-101.")

    if "COUPLING" in d and "RED" not in d:
        if len(mms) != 1:
            return unmatched("Equal coupling expected one size.")
        if mms[0] in S101:
            return hit("S-101", [mms[0]], ask3)
        return unmatched(f"φ{mms[0]} not listed on S-101.")

    if "REDUCER" in d:
        if len(mms) != 2:
            return unmatched("Reducer expected two sizes.")
        pair = (mms[0], mms[1]) if mms[0] >= mms[1] else (mms[1], mms[0])
        if pair in S102:
            return hit("S-102", list(pair), ask3)
        return unmatched(f"{fmt_phi(list(pair))} not listed on S-102. S-103 is a spigot reducer — ask only if you want fitting-reduced, not tube-reduced.")

    if "RED TEE" in d or "REDTEE" in c:
        return match_red_tee(mms, ask3)

    if d.strip() == "TEE" or c.endswith("-TEE"):
        if len(mms) == 1 or (len(mms) == 3 and mms[0] == mms[1] == mms[2]):
            mm = mms[0]
            if mm in T101:
                return hit("T-101", [mm, mm, mm], ask3)
            return unmatched(f"φ{mm} equal tee not on T-101.")
        return match_red_tee(mms, ask3)

    return unmatched("No factory family mapped for this APBS description.")


def match_red_tee(mms: list[int], ask3: str) -> dict:
    # US: often run x run x branch. Factory T-102: run x branch x run (1-2-3, 2 = branch).
    if len(mms) == 2:
        a, b = mms
        if a == b:
            if a in T101:
                return hit("T-101", [a, a, a], ask3)
        if a > b and (a, b, a) in T102:
            note = f"Interpreted {MM_TO_INCH.get(a,'?')}\" x {MM_TO_INCH.get(b,'?')}\" as equal-run reducing tee {MM_TO_INCH.get(a)}\" x {MM_TO_INCH.get(a)}\" x {MM_TO_INCH.get(b)}\" → factory 1×2×3 = {fmt_phi([a,b,a])}."
            return hit("T-102", [a, b, a], ask3, extra=note)
        if b > a and (a, b, a) in T103:
            return hit("T-103", [a, b, a], ask3, extra="Interpreted as bullhead (branch larger than run).")
        return unmatched(
            f"No T-102/T-103 for {fmt_phi(mms)}. Factory reduced-branch list does not include this combo."
        )

    if len(mms) != 3:
        return unmatched("Tee needs 1, 2, or 3 size parts.")

    a, b, c = mms
    # equal
    if a == b == c and a in T101:
        return hit("T-101", [a, a, a], ask3)
    # US run-run-branch (a==b, branch=c) → factory T-102 (a, c, a) if branch smaller
    if a == b and a > c and (a, c, a) in T102:
        extra = f"US {MM_TO_INCH.get(a)}\"×{MM_TO_INCH.get(a)}\"×{MM_TO_INCH.get(c)}\" = factory T-102 {fmt_phi([a,c,a])} (run×branch×run)."
        return hit("T-102", [a, c, a], ask3, extra=extra)
    # US 1/2 x 1/2 x 3/4 bullhead (branch larger) → T-103 15×22×15
    if a == b and c > a and (a, c, a) in T103:
        extra = f"US {MM_TO_INCH.get(a)}\"×{MM_TO_INCH.get(a)}\"×{MM_TO_INCH.get(c)}\" bullhead = factory T-103 {fmt_phi([a,c,a])}."
        return hit("T-103", [a, c, a], ask3, extra=extra)
    # US run-branch-run already factory order
    if a == c and a > b and (a, b, a) in T102:
        return hit("T-102", [a, b, a], ask3)
    # bullhead T-103 run-branch-run with branch larger
    if a == c and b > a and (a, b, a) in T103:
        return hit("T-103", [a, b, a], ask3)
    if a == c and b > a and (a, b, a) not in T103:
        # 1/2 x 3/4 x 1/2 style
        if (a, b, a) in T103:
            return hit("T-103", [a, b, a], ask3)
    # T-104 exact
    if (a, b, c) in T104:
        return hit("T-104", [a, b, c], ask3)
    # US 3/4 x 1/2 x 1/2
    if (a, b, c) in T104 or (a, c, b) in T104:
        tup = (a, b, c) if (a, b, c) in T104 else (a, c, b)
        return hit("T-104", list(tup), ask3)
    return unmatched(
        f"No T-101/102/103/104 listing for {fmt_phi(mms)}. Confirm orientation (run×run×branch vs factory 1×2×3)."
    )


def hit(pn: str, mms: list[int], ask3: str, extra: str = "") -> dict:
    name, zh, page, img = FAMILY[pn]
    notes = extra
    status = "MATCH"
    if ask3:
        status = "MATCH — confirm ASTM vs metric OD"
        notes = (notes + " " + ask3).strip()
    return {
        "factory_pn": pn,
        "factory_name": f"{name} / {zh}",
        "factory_size": fmt_phi(mms),
        "catalog_page": page,
        "photo": img,
        "status": status,
        "notes": notes,
    }


def unmatched(reason: str) -> dict:
    return {
        "factory_pn": "",
        "factory_name": "",
        "factory_size": "",
        "catalog_page": "",
        "photo": "",
        "status": "NOT IN THIS CATALOG",
        "notes": reason,
    }


def load_site():
    rows = []
    with SITE_CSV.open(newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            mat = (r.get("Material") or "").strip()
            if mat.lower() not in ("copper", "brass"):
                continue
            rows.append(r)
    return rows


def build_rows():
    out = []
    for r in load_site():
        code = r["Code"].strip()
        desc = r["Description"].strip()
        size = r["Size"].strip()
        parts = parse_inch_parts(size)
        mms = to_mm(parts) if parts else None
        m = match_fitting(code, desc, parts, mms)
        typ = "Type K" if "-K-" in code else ("Type L" if "-L-" in code else "")
        out.append(
            {
                "APBS_Code": code,
                "APBS_Description": desc,
                "APBS_Size": size if size.endswith('"') else size + '"',
                "Type_KL": typ,
                "APBS_Image": r.get("Image") or "",
                **m,
                "Qty": "",
                "FOB_USD": "",
                "Notes_to_factory": m["notes"],
            }
        )
    return out


HEADERS = [
    "Photo",
    "Factory_Item_No",
    "Factory_Name",
    "Factory_Size",
    "Catalog_Page",
    "Match",
    "APBS_Code",
    "APBS_Description",
    "APBS_Size",
    "Type_K_or_L",
    "Qty",
    "Unit_Price_FOB_USD",
    "Notes",
]


def write_csv(rows, path: Path):
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(
            f,
            fieldnames=[
                "Factory_Item_No",
                "Factory_Name",
                "Factory_Size",
                "Catalog_Page",
                "Match",
                "APBS_Code",
                "APBS_Description",
                "APBS_Size",
                "Type_K_or_L",
                "Qty",
                "Unit_Price_FOB_USD",
                "Notes",
            ],
        )
        w.writeheader()
        for r in rows:
            w.writerow(
                {
                    "Factory_Item_No": r["factory_pn"],
                    "Factory_Name": r["factory_name"],
                    "Factory_Size": r["factory_size"],
                    "Catalog_Page": r["catalog_page"],
                    "Match": r["status"],
                    "APBS_Code": r["APBS_Code"],
                    "APBS_Description": r["APBS_Description"],
                    "APBS_Size": r["APBS_Size"],
                    "Type_K_or_L": r["Type_KL"],
                    "Qty": "",
                    "Unit_Price_FOB_USD": "",
                    "Notes": r["notes"],
                }
            )


def thumb(src: Path, dest: Path, max_w=160, max_h=110):
    im = PILImage.open(src).convert("RGB")
    im.thumbnail((max_w, max_h), PILImage.Resampling.LANCZOS)
    dest.parent.mkdir(parents=True, exist_ok=True)
    im.save(dest, "PNG")
    return dest


def write_xlsx(rows):
    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    green = PatternFill("solid", fgColor="D9EAD3")
    yellow = PatternFill("solid", fgColor="FFF2CC")
    red = PatternFill("solid", fgColor="F4CCCC")
    thin = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )
    wrap = Alignment(wrap_text=True, vertical="center")

    # Cover
    cover = wb.active
    cover.title = "How_to_use"
    cover["A1"] = "Hailiang / AILIANG copper + brass RFQ"
    cover["A1"].font = Font(bold=True, size=16, color="1F4E79")
    blurb = [
        "Factory catalog excerpt: COPPER FITTING SERIES pp. 19–22 + BRASS FITTING SERIES pp. 23–25 (AILIANG / 海亮).",
        "This workbook is every copper SKU currently on allprobuildingsupplies.com (assets/products.csv). The site has 96 copper rows and 0 brass rows.",
        "Send the RFQ sheet. Qty and Unit_Price_FOB_USD are blank for them to fill. Pictures and item numbers are from THEIR catalog.",
        "Sizes on the catalog are φ mm (15/22/28/35/42/54/67/76/108…). Mapped to US CTS: 1/2\"=15, 3/4\"=22, 1\"=28, 1-1/4\"=35, 1-1/2\"=42, 2\"=54, 2-1/2\"=67, 3\"=76, 4\"=108.",
        "WARNING: 1/2\"–2-1/2\" OD is close to ASTM B88. 3\" ASTM OD is 79.375 mm (catalog φ76) and 4\" is 104.8 mm (catalog φ108). Those lines are MATCH — confirm ASTM vs metric.",
        "Type K vs Type L on our site is the TUBE wall. Wrot fittings are the same part (L-101, T-101, …). We still list both APBS codes so they quote one fitting that covers both if they agree.",
        "Pipe, PEX adapters, and stub-outs are not on these 7 pages. Keep those rows so they see what else we buy.",
        "Brass series is on the Brass_factory_only sheet — none of those SKUs are on our site yet.",
        "Ask: ASTM B16.22 / B88, NSF/cUPC if potable, FOB Ningbo, lead-free, carton qty, CBM.",
    ]
    cover["A2"] = "\n".join(f"• {b}" for b in blurb)
    cover["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    cover.row_dimensions[2].height = 180
    cover.column_dimensions["A"].width = 120

    # Size key
    sk = wb.create_sheet("Size_key")
    sk.append(["US nominal (CTS)", "ASTM B88 OD mm", "Hailiang catalog φ mm", "OK for US tube?"])
    for inch, mm in INCH_TO_MM.items():
        ok = "Yes (OD within ~0.9 mm)" if inch not in ("3", "4") else "ASK — metric φ ≠ ASTM OD"
        sk.append([inch + '"', ASTM_OD_MM[inch], mm, ok])
    for cell in sk[1]:
        cell.fill = header_fill
        cell.font = header_font
    for col in range(1, 5):
        sk.column_dimensions[get_column_letter(col)].width = 28

    # RFQ
    ws = wb.create_sheet("RFQ_send_this", 1)
    for col, h in enumerate(HEADERS, 1):
        cell = ws.cell(1, col, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:M{len(rows)+1}"

    tmp = Path("/tmp/hailiang_rfq_thumbs")
    tmp.mkdir(exist_ok=True)
    widths = [22, 16, 28, 16, 12, 28, 18, 26, 18, 12, 10, 18, 50]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    for i, r in enumerate(rows, 2):
        ws.row_dimensions[i].height = 78
        vals = [
            "",
            r["factory_pn"],
            r["factory_name"],
            r["factory_size"],
            r["catalog_page"] or "",
            r["status"],
            r["APBS_Code"],
            r["APBS_Description"],
            r["APBS_Size"],
            r["Type_KL"],
            "",
            "",
            r["notes"],
        ]
        for col, v in enumerate(vals, 1):
            cell = ws.cell(i, col, v)
            cell.border = thin
            cell.alignment = wrap
        st = r["status"]
        fill = None
        if st == "MATCH":
            fill = green
        elif st.startswith("MATCH"):
            fill = yellow
        else:
            fill = red
        ws.cell(i, 6).fill = fill

        img_path = None
        if r["photo"] and (CROPS / r["photo"]).exists():
            img_path = CROPS / r["photo"]
        elif r["APBS_Image"]:
            p = SITE_IMG / Path(r["APBS_Image"]).name
            if p.exists():
                img_path = p
        if img_path:
            th = thumb(img_path, tmp / f"r{i}.png")
            xl = XLImage(str(th))
            xl.width = 118
            xl.height = 72
            ws.add_image(xl, f"A{i}")

    # Catalog pages
    pg = wb.create_sheet("Catalog_pages")
    pg["A1"] = "Factory catalog pages used for this RFQ (excerpt pp. 19–25)"
    pg["A1"].font = Font(bold=True, size=14)
    row = 3
    for n, label in enumerate(
        [
            "19 Copper — couplings, reducers, 90° elbow",
            "20 Copper — reducing/45 elbow, tees, cap",
            "21 Copper — reducing tee T-102, U-bend, bushings, cross",
            "22 Copper — crossover, flush bushing, strap, female adapter",
            "23 Brass — male/female connectors, copper×thread",
            "24 Brass — unions, 90° F/M, female tee",
            "25 Brass — copper×thread tees, hangers, flange",
        ],
        19,
    ):
        pg.cell(row, 1, f"Page {n} — {label}")
        pg.row_dimensions[row].height = 18
        row += 1
        src = CROPS / f"page_{n}.png"
        if src.exists():
            th = thumb(src, tmp / f"page_{n}.png", 420, 600)
            xl = XLImage(str(th))
            xl.width = 320
            xl.height = 450
            pg.add_image(xl, f"A{row}")
            pg.row_dimensions[row].height = 340
        row += 1
    pg.column_dimensions["A"].width = 70

    # Brass factory-only
    br = wb.create_sheet("Brass_factory_only")
    br["A1"] = "Brass series is in this catalog (pp. 23–25) but All Pro’s live site has zero brass SKUs."
    br["A1"].font = Font(bold=True)
    br.append([])
    br.append(["Factory_Item_No", "Name", "Typical sizes", "Page", "Note"])
    brass_rows = [
        ("Smm-301", "Male Connector", '1/2"–2" M', 23, "Not on APBS site"),
        ("Smm-R302", "Reduce Male Connector", "reducing M×M", 23, "Not on APBS site"),
        ("Smc-303", "Male To Copper Connector", "φ15×1/2\" MC … φ54×2\" MC", 23, "Not on APBS site"),
        ("Sfc-304", "Female To Copper Connector", "φ15×1/2\" FC … φ54×2\" FC", 23, "Not on APBS site"),
        ("Sff-305", "Female Connector", '1/2"–2" F', 23, "Not on APBS site"),
        ("Sff-R306", "Reduce Female Connector", "reducing F×F", 23, "Not on APBS site"),
        ("Sfm-307", "Female & Male Connector", '1/2"–2" FM', 24, "Not on APBS site"),
        ("Sff-H308", "Removable Female Connector (union)", '1/2"–2" F', 24, "Not on APBS site"),
        ("Sfc-H309", "Removable Copper Connector", "φ×F union", 24, "Not on APBS site"),
        ("Smc-H310", "Removable Male To Copper", "φ×M union", 24, "Not on APBS site"),
        ("Lff-301", "90° Female Elbow", '1/2"–2" F', 24, "Not on APBS site"),
        ("Lfc-302", "90° Female To Copper Elbow", "φ×F", 24, "Not on APBS site"),
        ("Lmc-303", "90° Male To Copper Elbow", "φ×M", 24, "Not on APBS site"),
        ("Tff-301", "Female Tee", '1/2"–2" F', 24, "Not on APBS site"),
        ("Tfc-302", "Female To Copper Tee", "φ × F × φ", 25, "Not on APBS site"),
        ("Tfc-303", "Male Tee (copper×M)", "φ × M × φ", 25, "Not on APBS site"),
        ("TH", "Split Ring Hanger", "φ15–35", 25, "Not on APBS site"),
        ("BH", "Wall Ring", "φ15–35", 25, "Not on APBS site"),
        ("Sff-305 (flange)", "Copper liner + loose steel flange", "φ28–325", 25, "Same code as female connector — confirm with factory"),
    ]
    for row in brass_rows:
        br.append(list(row))
    for cell in br[3]:
        cell.fill = header_fill
        cell.font = header_font
    for col in range(1, 6):
        br.column_dimensions[get_column_letter(col)].width = 36

    wb.save(OUT_XLSX)


def main():
    rows = build_rows()
    write_csv(rows, OUT_CSV)
    write_xlsx(rows)
    n = len(rows)
    match = sum(1 for r in rows if r["status"] == "MATCH")
    ask = sum(1 for r in rows if r["status"].startswith("MATCH —"))
    miss = sum(1 for r in rows if r["status"].startswith("NOT"))
    print(f"APBS copper SKUs: {n}")
    print(f"MATCH: {match}  MATCH confirm ASTM: {ask}  NOT IN CATALOG: {miss}")
    print(f"Wrote {OUT_XLSX.name} and {OUT_CSV.name}")
    print()
    print(f"{'APBS':22} {'Size':20} {'Factory':8} {'φ':16} { 'Status'}")
    for r in rows:
        print(
            f"{r['APBS_Code'][:22]:22} {r['APBS_Size'][:20]:20} {r['factory_pn'] or '—':8} "
            f"{r['factory_size'][:16]:16} {r['status']}"
        )


if __name__ == "__main__":
    main()
