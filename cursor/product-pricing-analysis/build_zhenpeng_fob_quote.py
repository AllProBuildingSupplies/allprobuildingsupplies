#!/usr/bin/env python3
"""Zhenpeng FOB quote 2026-09-05 vs NJPD/Gator sell, Tommur DDP, King Smart."""

from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None

OUT = Path(__file__).resolve().parent
DUTY = 0.428
# 163 cartons of fittings ≈ 5 CBM. $3,300/40' / 67.7 CBM ≈ $48.74/CBM → ~$0.004/pc.
# Use $0.01/pc so landed is not understated.
FREIGHT_PC = 0.01

# Quote to Baruch Grossman, All Pro Building Supplies LLC, 2026/9/5, FOB Ningbo.
# PPSU Solvay, ASTM F2159, cUPC + NSF/ANSI 61. Wholesale pack. 55,100 pcs / $16,686.15.
# gator, zp_pn, desc, size, qty, fob, njpd_sell, tommur_fob, tommur_ddp, ks_fob_low, ks_usa
ROWS = [
    ("PPCP0012", "B1212CP", '1/2" PEX coupling', '1/2"', 3000, 0.107, 0.2304, 0.1576, 0.25, None, None),
    ("PPCP0034", "B3434CP", '3/4" PEX coupling', '3/4"', 4500, 0.207, 0.3795, 0.2406, 0.40, None, None),
    ("PPCP0100", "B0505CP", '1" PEX coupling', '1"', 1200, 0.440, 0.7695, 0.5090, 0.81, None, None),
    ("PPRC3412", "B3412CP", '3/4"x1/2" PEX coupling', '3/4x1/2"', 2000, 0.180, 0.3313, 0.2101, 0.34, None, None),
    ("PPRC1034", "B0534CP", '1"x3/4" PEX coupling', '1x3/4"', 1800, 0.330, 0.7138, None, None, None, None),
    ("PPLN0012", "B121290", '1/2" PEX elbow', '1/2"', 9000, 0.172, 0.3313, 0.2933, 0.49, None, None),
    ("PPLN0034", "B343490", '3/4" PEX elbow', '3/4"', 4500, 0.353, 0.6581, 0.4000, 0.83, 0.51, 0.83),
    ("PPLN0100", "B050590", '1" PEX elbow', '1"', 2700, 0.592, 1.4607, None, None, None, None),
    ("PPPL0012", "B12PLG", '1/2" PEX plug', '1/2"', 6000, 0.094, 0.2078, None, None, None, None),
    ("PPPL0034", "B34PLG", '3/4" PEX plug', '3/4"', 3000, 0.156, 0.3629, None, None, None, None),
    ("PPPL0100", "B05PLG", '1" PEX plug', '1"', 1500, 0.268, 0.5903, None, None, 0.47, 0.68),
    ("PPTE0012", "B12T", '1/2" PEX tee', '1/2"', 1200, 0.237, 0.4488, 0.2955, 0.47, 0.38, 0.49),
    ("PPTE0034", "B34T", '3/4" PEX tee', '3/4"', 1800, 0.489, 0.9231, 0.5254, 0.84, None, None),
    ("PPTE0100", "B05T", '1" PEX tee', '1"', 600, 0.938, 2.0344, 0.9060, 1.45, None, None),
    ("PPRT1213", "B121234T", '1/2"x1/2"x3/4" PEX tee', '1/2x1/2x3/4"', 900, 0.346, 0.6445, None, None, None, None),
    ("PPRT3411", "B341212T", '3/4"x1/2"x1/2" PEX tee', '3/4x1/2x1/2"', 2700, 0.383, 0.5963, None, None, None, None),
    ("PPRT3413", "B341234T", '3/4"x1/2"x3/4" PEX tee', '3/4x1/2x3/4"', 2250, 0.439, 0.7891, None, None, 0.65, 1.04),
    ("PPRT3431", "B343412T", '3/4"x3/4"x1/2" PEX tee', '3/4x3/4x1/2"', 3000, 0.455, 0.7439, 0.4900, 0.78, None, None),
    ("PPRT3410", "B343405T", '3/4"x3/4"x1" PEX tee', '3/4x3/4x1"', 450, 0.590, 1.4366, None, None, None, None),
    ("PPRT1033", "B053434T", '1"x3/4"x3/4" PEX tee', '1x3/4x3/4"', 900, 0.662, 1.6745, None, None, None, None),
    ("PPRT1341", "B053405T", '1"x3/4"x1" PEX tee', '1x3/4x1"', 600, 0.737, 2.0525, None, None, None, None),
    ("PPRT1112", "B050512T", '1"x1"x1/2" PEX tee', '1x1x1/2"', 900, 0.682, 1.3929, None, None, None, None),
    ("PPRT1134", "B050534T", '1"x1"x3/4" PEX tee', '1x1x3/4"', 600, 0.733, 1.5782, 0.8821, 1.41, None, None),
]


def landed(fob):
    return round(fob * (1 + DUTY) + FREIGHT_PC, 4)


def pct_vs(a, b):
    if a is None or b in (None, 0):
        return None
    return round((a - b) / b * 100, 1)


def margin(sell, cost):
    if sell in (None, 0) or cost is None:
        return None
    return round((sell - cost) / sell * 100, 1)


def band(m):
    if m is None:
        return ""
    if m < 0:
        return "UNDERWATER"
    if m < 20:
        return "thin (<20%)"
    if m < 50:
        return "OK vs NJPD, below 50% target"
    if m < 75:
        return "hits 50% target"
    return "hits 75% target"


def better(zp_l, tddp):
    if tddp is None:
        return "no Tommur DDP — Zhenpeng covers a hole"
    if zp_l < tddp:
        return "Zhenpeng landed cheaper"
    return "Tommur DDP cheaper"


def csv_escape(v):
    s = "" if v is None else str(v)
    if any(c in s for c in ',\"\n'):
        return '"' + s.replace('"', '""') + '"'
    return s


HEADERS = [
    "Gator",
    "Zhenpeng_PN",
    "Item",
    "Size",
    "Quote_qty",
    "Zhenpeng_FOB_USD",
    "Zhenpeng_landed_USD",
    "NJPD_sell_USD",
    "Cost_for_50pct",
    "Cost_for_75pct",
    "Margin_on_Zhenpeng_landed_pct",
    "Band",
    "Tommur_FOB_USD",
    "Tommur_DDP_USD",
    "ZP_FOB_vs_Tommur_FOB_pct",
    "ZP_landed_vs_Tommur_DDP_pct",
    "Better_buy",
    "Margin_on_Tommur_DDP_pct",
    "KingSmart_FOB_low_USD",
    "KingSmart_USA_USD",
    "Line_FOB_USD",
    "Line_landed_USD",
    "Line_NJPD_sell_USD",
    "Line_profit_vs_NJPD_USD",
]


def rows():
    out = []
    for gator, pn, item, size, qty, fob, sell, tfob, tddp, ks_fob, ks_usa in ROWS:
        l = landed(fob)
        out.append(
            {
                "Gator": gator,
                "Zhenpeng_PN": pn,
                "Item": item,
                "Size": size,
                "Quote_qty": qty,
                "Zhenpeng_FOB_USD": fob,
                "Zhenpeng_landed_USD": l,
                "NJPD_sell_USD": sell,
                "Cost_for_50pct": round(sell * 0.50, 4),
                "Cost_for_75pct": round(sell * 0.25, 4),
                "Margin_on_Zhenpeng_landed_pct": margin(sell, l),
                "Band": band(margin(sell, l)),
                "Tommur_FOB_USD": tfob,
                "Tommur_DDP_USD": tddp,
                "ZP_FOB_vs_Tommur_FOB_pct": pct_vs(fob, tfob),
                "ZP_landed_vs_Tommur_DDP_pct": pct_vs(l, tddp),
                "Better_buy": better(l, tddp),
                "Margin_on_Tommur_DDP_pct": margin(sell, tddp),
                "KingSmart_FOB_low_USD": ks_fob,
                "KingSmart_USA_USD": ks_usa,
                "Line_FOB_USD": round(fob * qty, 2),
                "Line_landed_USD": round(l * qty, 2),
                "Line_NJPD_sell_USD": round(sell * qty, 2),
                "Line_profit_vs_NJPD_USD": round((sell - l) * qty, 2),
            }
        )
    return out


def write_csv(data, path):
    lines = [",".join(HEADERS)]
    for r in data:
        lines.append(",".join(csv_escape(r[h]) for h in HEADERS))
    path.write_text("\n".join(lines) + "\n")


def write_xlsx(data, path):
    if Workbook is None:
        return
    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    green = PatternFill("solid", fgColor="D9EAD3")
    red = PatternFill("solid", fgColor="F4CCCC")
    yellow = PatternFill("solid", fgColor="FFF2CC")
    thin = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )
    money_cols = {
        "Zhenpeng_FOB_USD",
        "Zhenpeng_landed_USD",
        "NJPD_sell_USD",
        "Cost_for_50pct",
        "Cost_for_75pct",
        "Tommur_FOB_USD",
        "Tommur_DDP_USD",
        "KingSmart_FOB_low_USD",
        "KingSmart_USA_USD",
        "Line_FOB_USD",
        "Line_landed_USD",
        "Line_NJPD_sell_USD",
        "Line_profit_vs_NJPD_USD",
    }
    pct_cols = {
        "Margin_on_Zhenpeng_landed_pct",
        "ZP_FOB_vs_Tommur_FOB_pct",
        "ZP_landed_vs_Tommur_DDP_pct",
        "Margin_on_Tommur_DDP_pct",
    }
    ws = wb.active
    ws.title = "Fittings"
    for col, h in enumerate(HEADERS, 1):
        c = ws.cell(1, col, h)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, r in enumerate(data, 2):
        for col, h in enumerate(HEADERS, 1):
            v = r[h]
            cell = ws.cell(i, col, v)
            cell.border = thin
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            if h in money_cols and isinstance(v, (int, float)):
                cell.number_format = '"$"#,##0.0000' if abs(v) < 20 else '"$"#,##0.00'
            if h in pct_cols and isinstance(v, (int, float)):
                cell.number_format = '0.0"%"'
                if h == "Margin_on_Zhenpeng_landed_pct":
                    cell.fill = red if v < 0 else (yellow if v < 50 else green)
                if h in ("ZP_FOB_vs_Tommur_FOB_pct", "ZP_landed_vs_Tommur_DDP_pct") and v < 0:
                    cell.fill = green
                if h in ("ZP_FOB_vs_Tommur_FOB_pct", "ZP_landed_vs_Tommur_DDP_pct") and v > 0:
                    cell.fill = red
            if h == "Better_buy" and v and v.startswith("Zhenpeng"):
                cell.fill = green
            if h == "Band" and v == "UNDERWATER":
                cell.fill = red
            if h == "Band" and v == "OK vs NJPD, below 50% target":
                cell.fill = yellow
            if h == "Band" and "hits" in (v or ""):
                cell.fill = green
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(data)+1}"
    ws.row_dimensions[1].height = 32
    widths = [12, 12, 28, 16, 12, 14, 16, 14, 14, 14, 16, 28, 14, 14, 16, 16, 28, 16, 16, 14, 14, 14, 16, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    info = wb.create_sheet("Source")
    facts = [
        ("Seller", "Ningbo Zhenpeng Plumbing Fittings Co., Ltd. (ZHNP)"),
        ("Buyer", "Baruch Grossman / All Pro Building Supplies LLC"),
        ("Date", "2026/9/5"),
        ("Incoterm", "FOB Ningbo"),
        ("Spec", "ASTM F2159 PPSU Solvay. cUPC + NSF/ANSI 61 claimed. Wholesale pack."),
        ("Totals", "55,100 pcs / $16,686.15 FOB / 163 cartons / 1,915 bags"),
        ("Pay", "30% deposit, 70% T/T before ship, 30 days, FX ±1% clause"),
        ("HS", "3917400000"),
        ("Landed", "FOB × 1.428 (5.3% MFN + 25% Sec 301 + 12.5% FLIP) + $0.01 freight/pc"),
        ("Published vs quoted", "B121290 1/2\" elbow was listed $0.10 FOB MOQ 3000; this quote is $0.172 (+72%)"),
        ("Not on quote", "No 20ft PEX-B pipe. No copper F1807 crimp rings."),
        ("King Smart", "USA warehouse still not a production buy. Monday FOB CN still pending."),
    ]
    info["A1"] = "Field"
    info["B1"] = "Value"
    info["A1"].fill = header_fill
    info["B1"].fill = header_fill
    info["A1"].font = header_font
    info["B1"].font = header_font
    for i, (k, v) in enumerate(facts, 2):
        info.cell(i, 1, k)
        info.cell(i, 2, v)
        info.cell(i, 2).alignment = Alignment(wrap_text=True)
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 110
    wb.save(path)


def main():
    data = rows()
    write_csv(data, OUT / "Zhenpeng_FOB_vs_NJPD.csv")
    write_xlsx(data, OUT / "Zhenpeng_FOB_vs_NJPD.xlsx")
    fob_tot = sum(r["Line_FOB_USD"] for r in data)
    land_tot = sum(r["Line_landed_USD"] for r in data)
    sell_tot = sum(r["Line_NJPD_sell_USD"] for r in data)
    profit = sell_tot - land_tot
    m_order = (profit / sell_tot) * 100 if sell_tot else 0
    underwater = sum(1 for r in data if r["Margin_on_Zhenpeng_landed_pct"] < 0)
    hit50 = sum(1 for r in data if r["Margin_on_Zhenpeng_landed_pct"] >= 50)
    vs_ddp = [r for r in data if r["Tommur_DDP_USD"] is not None]
    cheaper = sum(1 for r in vs_ddp if r["Zhenpeng_landed_USD"] < r["Tommur_DDP_USD"])
    print(f"FOB ${fob_tot:,.2f}  landed ~${land_tot:,.2f}  NJPD sell-through ${sell_tot:,.2f}")
    print(f"Order margin at NJPD prices: {m_order:.1f}%  profit ${profit:,.0f}")
    print(f"Underwater: {underwater}/{len(data)}  hits 50%: {hit50}/{len(data)}")
    print(f"Beats Tommur DDP: {cheaper}/{len(vs_ddp)}")
    print()
    print(f"{'Item':28} {'FOB':>7} {'Land':>7} {'Sell':>7} {'M%':>6} {'T-DDP':>7} vsDDP  { '50%?':5}")
    for r in data:
        t = "" if r["Tommur_DDP_USD"] is None else f"{r['Tommur_DDP_USD']:.2f}"
        vd = "" if r["ZP_landed_vs_Tommur_DDP_pct"] is None else f"{r['ZP_landed_vs_Tommur_DDP_pct']:+.0f}%"
        hit = "YES" if r["Margin_on_Zhenpeng_landed_pct"] >= 50 else "no"
        print(
            f"{r['Item'][:28]:28} {r['Zhenpeng_FOB_USD']:7.3f} {r['Zhenpeng_landed_USD']:7.3f} "
            f"{r['NJPD_sell_USD']:7.3f} {r['Margin_on_Zhenpeng_landed_pct']:5.1f}% {t:>7} {vd:>5}  {hit}"
        )


if __name__ == "__main__":
    main()
