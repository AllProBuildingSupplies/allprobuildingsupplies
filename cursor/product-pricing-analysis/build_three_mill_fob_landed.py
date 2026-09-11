#!/usr/bin/env python3
"""Nine-column Palconn / Tommur / Zhenpeng FOB + landed vs APBS sell. Internal.

Union of every priced quote we have from the three mills. PEX pipe and fittings
are 1/2\", 3/4\", 1\" only. Blank cell = that mill did not quote a number.
Do not invent prices. Buyer decides what to order.
"""

from __future__ import annotations

import csv
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE / "palconn"))
from build_palconn_quote_compare import (  # noqa: E402
    CBM_40,
    DUTY,
    OCEAN,
    QUOTE_10,
    QUOTE_8,
    ZHENPENG,
    canon_size,
    load_quote,
    num,
    ocean_per_unit,
)

FACTORY = HERE / "Factory_Order_PVC_PEX_45HQ.csv"
TRACKER = HERE / "Tommur_Cost_Margin_Tracker.xlsx"
SITE = ROOT / "assets" / "products.csv"
OUT_XLSX = HERE / "Three_Mill_FOB_Landed.xlsx"
OUT_CSV = HERE / "Three_Mill_FOB_Landed.csv"

CBM_45HQ = 86.0
ZHEN_OCEAN_PC = OCEAN * 2.3 / CBM_45HQ / 55_100

PEX_OK = {"1/2", "3/4", "1"}
PEX_BAD = {"1-1/4", "1-1/2", "2", "2-1/2"}

HEADER = PatternFill("solid", fgColor="C00000")
WHITE = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
GREEN = PatternFill("solid", fgColor="C6EFCE")
WRAP = Alignment(wrap_text=True, vertical="center")
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)

COLS = [
    ("item", "APBS item", 52),
    ("size", "Size", 18),
    ("pal_fob", "Palconn FOB", 14),
    ("tom_fob", "Tommur FOB", 14),
    ("zhen_fob", "Zhenpeng FOB", 14),
    ("pal_landed", "Palconn est landed", 18),
    ("tom_landed", "Tommur DDP / est landed", 22),
    ("zhen_landed", "Zhenpeng est landed", 18),
    ("apbs_sell", "APBS selling price", 16),
]

DESC_TO_CODE = (
    ("CPVC ASTM 2846 PIPE SDR-13.5", "CPVC-PIPE-SDR13.5"),
    ("CPVC ASTM 2846 PIPE SDR-11", "CPVC-PIPE-SDR11"),
    ("ASTM CPVC SCH80 PIPE", "CPVC-PIPE-SCH80"),
    ("ASTM D1785 SCH40 PVC", "PVC-PIPE-SOLID"),
    ("ASTM F891 PVC FOAM CORE", "PVC-PIPE-FOAM"),
    ("ASTM UPVC SCH80 PIPE", "PVC-PIPE-SCH80"),
    ("45° ELL", "PVC-45ELLSOC"),
    ("CLOSET FLANGE W/ADJUSTABLE METAL RING", "PVC-CLSTFLNG-ADJ"),
)

CPVC_FITTING_CODE = {
    "COUPLING": "CPVC-CPLNG",
    "TEE": "CPVC-TEE",
}


def money(v, nd=4):
    if v is None:
        return None
    return round(float(v), nd)


def size_key(raw) -> str:
    s = str(raw or "")
    s = (
        s.replace("1½", "1-1/2")
        .replace("1¼", "1-1/4")
        .replace("2½", "2-1/2")
        .replace("½", "1/2")
        .replace("¼", "1/4")
        .replace("¾", "3/4")
    )
    s = re.sub(r"[（(].*$", "", s)
    return canon_size(s)


def size_tokens(size: str) -> list[str]:
    s = size_key(size)
    if not s:
        return []
    return [p for p in re.split(r"x", s) if p]


def is_pex(code: str, family: str | None = None) -> bool:
    blob = f"{code} {family or ''}".upper()
    return blob.startswith("PEX") or " F1807" in blob or " F2159" in blob or "PPSU" in blob


def pex_allowed(code: str, size: str, family: str | None = None) -> bool:
    if not is_pex(code, family):
        return True
    toks = size_tokens(size)
    if not toks:
        return False
    return all(t in PEX_OK for t in toks)


def fmt_size(raw) -> str:
    c = size_key(raw)
    if not c:
        return str(raw or "").strip()
    if "x" in c:
        return '" x '.join(c.split("x")) + '"'
    return c + '"'


def item_label(code, desc, unit=None) -> str:
    code = (code or "").strip()
    desc = (desc or "").strip()
    extra = ""
    dlow = desc.lower()
    if unit == "20 ft stick" or "20 ft stick" in dlow:
        extra = " (20 ft stick)"
    elif "100 ft coil" in dlow:
        extra = " (100 ft coil)"
    if code and desc:
        base = f"{code} — {desc}"
    else:
        base = code or desc
    if extra and extra.strip(" ()") not in dlow:
        return base + extra
    return base


def palconn_landed(row) -> float | None:
    fob = row.get("fob")
    if fob is None:
        return None
    ocean = ocean_per_unit(row) or 0.0
    return money(fob * DUTY + ocean)


def zhen_landed(fob: float | None) -> float | None:
    if fob is None:
        return None
    return money(fob * DUTY + ZHEN_OCEAN_PC)


def infer_code(code: str, desc: str) -> str:
    code = (code or "").strip()
    if code:
        return code
    d = (desc or "").strip().upper()
    if d in CPVC_FITTING_CODE:
        return CPVC_FITTING_CODE[d]
    for needle, mapped in DESC_TO_CODE:
        if needle.upper() in d:
            return mapped
    return ""


def load_tommur() -> dict:
    """Tracker (full Tommur book) + 45HQ factory CSV. Prefer a real FOB/DDP number from either."""
    out: dict[tuple[str, str], dict] = {}

    def put(code, size, rec):
        if not code or not size:
            return
        if rec.get("fob") == 0 and rec.get("ddp") in (None, 0) and rec.get("est") in (None, 0):
            return
        k = (code, size)
        prev = out.get(k)
        if prev is None:
            out[k] = rec
            return
        for field in ("fob", "ddp", "est", "desc", "unit", "size_raw"):
            if prev.get(field) in (None, "") and rec.get(field) not in (None, ""):
                prev[field] = rec[field]
        if rec.get("fob") is not None and prev.get("fob") is None:
            prev["fob"] = rec["fob"]
        if rec.get("ddp") is not None and prev.get("ddp") is None:
            prev["ddp"] = rec["ddp"]
        if rec.get("est") is not None and prev.get("est") is None:
            prev["est"] = rec["est"]

    wb = load_workbook(TRACKER, data_only=True)
    ws = wb["Cost_Margin_Master"]
    headers = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        r = dict(zip(headers, row))
        mat = str(r.get("Material") or "")
        if mat not in ("PVC", "CPVC", "PEX"):
            continue
        desc = str(r.get("Description") or "").strip()
        code = infer_code(str(r.get("APBS_Item_Code") or ""), desc)
        size = size_key(r.get("Size"))
        rec = {
            "code": code,
            "desc": desc,
            "size_raw": r.get("Size"),
            "unit": str(r.get("Sell_Unit") or "").strip(),
            "fob": num(r.get("FOB_USD")),
            "ddp": num(r.get("DDP_Current")),
            "est": num(r.get("Est_Landed_per_Pc_45HQ")),
            "per_ft": mat in ("PVC", "CPVC") and "PIPE" in (code + " " + desc).upper()
            and "FITTING" not in desc.upper()
            and "INCREASER" not in desc.upper()
            and "REDUCER" not in desc.upper()
            and "ELL" not in desc.upper(),
        }
        if rec["fob"] is None and rec["ddp"] is None and rec["est"] is None:
            continue
        put(code, size, rec)
    wb.close()

    with FACTORY.open(newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            code = (r.get("APBS_Code") or "").strip()
            size = size_key(r.get("Size") or "")
            if not code:
                continue
            rec = {
                "code": code,
                "desc": (r.get("Description") or "").strip(),
                "size_raw": r.get("Size"),
                "unit": (r.get("Unit") or "").strip(),
                "fob": num(r.get("FOB_USD")),
                "ddp": num(r.get("DDP_USD")),
                "est": num(r.get("Est_Landed_45HQ_USD")),
                "per_ft": code.startswith("PEX-B PIPE") or code.startswith("PVC-PIPE") or code.startswith("CPVC-PIPE"),
            }
            put(code, size, rec)

    for rec in out.values():
        rec["landed"] = rec["ddp"] if rec.get("ddp") is not None else rec.get("est")
        if rec["landed"] is None and rec.get("fob") is not None:
            rec["landed"] = money(rec["fob"] * DUTY)
    return out


def load_site() -> tuple[dict, dict]:
    prices, meta = {}, {}
    with SITE.open(newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            code = (r.get("Code") or "").strip()
            size = size_key(r.get("Size") or "")
            if not code:
                continue
            prices[(code, size)] = num(r.get("Price"))
            meta[(code, size)] = {
                "desc": (r.get("Description") or "").strip(),
                "material": (r.get("Material") or "").strip(),
            }
    return prices, meta


def palconn_map() -> dict:
    """10 Sep wins on overlap; 8 Sep keeps DWV (and anything else) she later omitted."""
    m = {}
    for src, recs in (("8", load_quote(QUOTE_8)), ("10", load_quote(QUOTE_10))):
        for r in recs:
            if not pex_allowed(str(r.get("code") or ""), str(r.get("size") or ""), r.get("family")):
                continue
            code = (r.get("code") or "").strip()
            size = size_key(r.get("size"))
            desc = str(r.get("desc") or "")
            variant = ""
            if "100 ft coil" in desc.lower():
                variant = "coil"
            elif "20 ft stick" in desc.lower() or r.get("unit") == "20 ft stick":
                variant = "stick"
            k = (code, size, variant)
            if src == "8" and k in m and m[k].get("fob") is not None:
                continue
            if src == "10" or r.get("fob") is not None or k not in m:
                row = dict(r)
                row["_variant"] = variant
                m[k] = row
    return m


def family_of(code: str, pal_family: str | None) -> str:
    if pal_family:
        return pal_family
    if code.startswith("PEX-F1807") or code in ("PEX-WMBOX", "PEX-WMBOX-WHA", "PEX-ICEBOX"):
        return "PEX F1807 / boxes"
    if code.startswith("PEX"):
        return "PEX"
    if code.startswith("PVC-PIPE"):
        return "PVC pipe"
    if code.startswith("CPVC"):
        return "CPVC"
    return "PVC DWV"


def tommur_on_palconn_row(tr: dict, pr: dict | None) -> tuple[float | None, float | None]:
    """Palconn PVC pipe is a 20 ft stick. Tommur book is $/ft — scale those cells ×20."""
    fob, landed = tr.get("fob"), tr.get("landed")
    if not pr:
        return fob, landed
    unit = str(pr.get("unit") or "")
    code = str(pr.get("code") or "")
    stick = unit == "20 ft stick" or pr.get("_variant") == "stick"
    if stick and code.startswith("PVC-PIPE"):
        if fob is not None:
            fob = money(fob * 20)
        if landed is not None:
            landed = money(landed * 20)
    return fob, landed


def build_rows() -> list[dict]:
    pal = palconn_map()
    tom = {k: v for k, v in load_tommur().items() if pex_allowed(k[0], k[1])}
    site, site_meta = load_site()
    zhen = {k: v for k, v in ZHENPENG.items() if pex_allowed(k[0], k[1])}

    pal_sizes = {(c, s) for (c, s, _v) in pal}
    keys = set(pal.keys())
    for code, size in list(tom) + list(zhen):
        if (code, size) not in pal_sizes:
            keys.add((code, size, ""))

    rows = []
    for code, size, variant in sorted(keys, key=lambda x: (x[0], x[1], x[2])):
        if not pex_allowed(code, size):
            continue
        pr = pal.get((code, size, variant))
        if pr is None and variant == "":
            pr = pal.get((code, size, "stick")) or pal.get((code, size, "coil"))
        tr = tom.get((code, size), {})
        sell = site.get((code, size))
        meta = site_meta.get((code, size), {})

        pal_fob = money(pr.get("fob")) if pr else None
        pal_l = palconn_landed(pr) if pr else None
        tom_fob, tom_l = tommur_on_palconn_row(tr, pr)
        tom_fob = money(tom_fob) if tom_fob is not None else None
        tom_l = money(tom_l) if tom_l is not None else None
        zhen_fob = zhen.get((code, size))
        zhen_l = zhen_landed(zhen_fob)

        if pal_fob is None and tom_fob is None and zhen_fob is None and pal_l is None and tom_l is None:
            continue

        desc = (pr or {}).get("desc") or tr.get("desc") or meta.get("desc") or ""
        unit = (pr or {}).get("unit") or tr.get("unit")
        size_disp = fmt_size((pr or {}).get("size") or tr.get("size_raw") or size)
        item = item_label(code, desc, unit)

        rows.append(
            {
                "family": family_of(code, (pr or {}).get("family")),
                "code": code,
                "item": item,
                "size": size_disp,
                "pal_fob": pal_fob,
                "tom_fob": tom_fob,
                "zhen_fob": zhen_fob,
                "pal_landed": pal_l,
                "tom_landed": tom_l,
                "zhen_landed": zhen_l,
                "apbs_sell": sell,
            }
        )

    order = {
        "PEX-B pipe": 0,
        "PEX": 1,
        "PEX F2159 PPSU": 2,
        "PEX F1807 brass": 3,
        "PEX F1807 / boxes": 3,
        "PEX outlet boxes": 4,
        "PVC Sch40 pipe": 5,
        "PVC foam-core pipe": 6,
        "PVC pipe": 5,
        "PVC DWV fittings": 7,
        "PVC DWV": 7,
        "CPVC": 8,
        "CPVC SDR-11 (if manufactured)": 8,
        "CPVC D2846 fittings (if manufactured)": 8,
    }
    rows.sort(key=lambda r: (order.get(r["family"], 9), r["code"], r["size"], r["item"]))
    return rows


def cheapest(vals):
    nums = [v for v in vals if isinstance(v, (int, float))]
    if not nums:
        return None
    return min(nums)


def write_xlsx(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Three mills"
    for i, (_, title, width) in enumerate(COLS, 1):
        c = ws.cell(1, i, title)
        c.fill = HEADER
        c.font = WHITE
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{len(rows)+1}"

    for ri, row in enumerate(rows, 2):
        fob_min = cheapest([row["pal_fob"], row["tom_fob"], row["zhen_fob"]])
        for ci, (k, _, _) in enumerate(COLS, 1):
            v = row.get(k)
            cell = ws.cell(ri, ci, v)
            cell.alignment = WRAP
            cell.border = THIN
            cell.font = Font(name="Calibri", size=10)
            if k in (
                "pal_fob",
                "tom_fob",
                "zhen_fob",
                "pal_landed",
                "tom_landed",
                "zhen_landed",
                "apbs_sell",
            ) and isinstance(v, (int, float)):
                cell.number_format = "0.00" if abs(v) >= 1 else "0.0000"
                if k.endswith("_fob") and fob_min is not None and abs(v - fob_min) < 1e-9:
                    cell.fill = GREEN
        ws.row_dimensions[ri].height = 22

    notes = wb.create_sheet("Notes")
    lines = [
        "INTERNAL — do not forward to any mill.",
        "",
        "One list: every priced quote we have from Palconn, Tommur, and Zhenpeng. Blank = that mill did not quote a number. You decide what to order.",
        "PEX pipe and fittings on this sheet are 1/2\", 3/4\", and 1\" only. 1-1/4 / 1-1/2 / 2\" PEX is off the list.",
        "Green = lowest FOB among mills that quoted that row.",
        "",
        "Palconn FOB = 10 Sep sheet; 8 Sep FOB kept on DWV she dropped off the 10 Sep return. Palconn est landed = FOB × 1.43 + $10,000 / 67.7 CBM × CBM per selling unit (ocean $0 when she did not give CBM).",
        "Tommur FOB = yellow FOB from Tommur_Cost_Margin_Tracker + Factory_Order_PVC_PEX_45HQ. Tommur landed = mill DDP if present, else Est_Landed_45HQ, else FOB × 1.43.",
        "Palconn PVC pipe is a 20 ft stick. Tommur PVC/CPVC pipe in the book is $/ft — on Palconn stick rows Tommur FOB and DDP/landed are ×20 so the cells sit next to Palconn. Tommur-only pipe sizes stay $/ft.",
        "Zhenpeng FOB = 5 Sep 2026 REV PI (10 SKUs). Est landed = FOB × 1.43 + $10,000 × 2.3 CBM / 86 CBM / 55,100 pcs. Plugs / extra tees were never on the REV PI — left blank.",
        "APBS selling price = assets/products.csv. Blank = not on the site. PVC 1-1/2\" solid on the site looks $/ft; 2\" / 3\" / 4\" look like stick prices.",
        "1\" F2159 female adapter: Palconn returned // — no number to fill.",
        "",
        f"Ocean constants: $10,000 / 40ft ({CBM_40} CBM) and same $10,000 / 45HQ ({CBM_45HQ} CBM). Duty stack 1.43.",
        "Rebuild: python3 cursor/product-pricing-analysis/build_three_mill_fob_landed.py",
    ]
    notes["A1"] = "Three-mill FOB / landed"
    notes["A1"].font = Font(bold=True, size=14, name="Calibri")
    for i, line in enumerate(lines, 3):
        notes.cell(i, 1, line).font = Font(name="Calibri", size=11)
    notes.column_dimensions["A"].width = 140
    wb.save(OUT_XLSX)


def assert_complete(rows):
    """Fail the rebuild if a known mill number did not land on the sheet."""
    by = {(r["code"], size_key(r["size"]), "coil" if "coil" in r["item"].lower() else "stick" if "stick" in r["item"].lower() else ""): r for r in rows}
    by_cs = {}
    for r in rows:
        by_cs.setdefault((r["code"], size_key(r["size"])), r)

    over = [r for r in rows if is_pex(r["code"]) and not pex_allowed(r["code"], r["size"])]
    if over:
        raise SystemExit(f"PEX oversize still on sheet: {[(r['code'], r['size']) for r in over]}")

    pal = palconn_map()
    missing_pal = []
    for (code, size, var), pr in pal.items():
        if pr.get("fob") is None:
            continue
        hit = None
        if var:
            hit = by.get((code, size, var))
        if hit is None:
            hit = by_cs.get((code, size))
        if hit is None or hit.get("pal_fob") is None:
            missing_pal.append((code, size, var, pr.get("fob")))
        elif hit.get("pal_landed") is None:
            missing_pal.append((code, size, var, "FOB but no landed"))
    if missing_pal:
        raise SystemExit(f"Palconn priced rows missing from sheet: {missing_pal[:12]}")

    for k, fob in ZHENPENG.items():
        if not pex_allowed(k[0], k[1]):
            continue
        hit = by_cs.get(k)
        if hit is None or hit.get("zhen_fob") is None or hit.get("zhen_landed") is None:
            raise SystemExit(f"Zhenpeng missing {k} {fob}")

    cpvc = by_cs.get(("CPVC-PIPE-SCH80", "1/2"))
    if not cpvc or cpvc.get("tom_fob") is None:
        raise SystemExit("Tommur CPVC-PIPE-SCH80 1/2 FOB missing")
    solid = None
    for r in rows:
        if r["code"] == "PVC-PIPE-SOLID" and size_key(r["size"]) == "2" and "stick" in r["item"].lower():
            solid = r
            break
    if not solid or solid.get("tom_fob") is None:
        raise SystemExit("Tommur FOB missing on Palconn 2\" PVC stick row")
    for r in rows:
        if r["code"] == "PEX-B PIPE" and size_key(r["size"]) in PEX_BAD:
            raise SystemExit(f"oversize PEX pipe {r['size']}")


def main():
    rows = build_rows()
    assert_complete(rows)
    write_xlsx(rows)
    with OUT_CSV.open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=[k for k, _, _ in COLS], extrasaction="ignore")
        w.writeheader()
        for row in rows:
            w.writerow({k: row.get(k, "") for k, _, _ in COLS})
    n = len(rows)
    pal = sum(1 for r in rows if r["pal_fob"] is not None)
    tom = sum(1 for r in rows if r["tom_fob"] is not None)
    zhen = sum(1 for r in rows if r["zhen_fob"] is not None)
    pal_l = sum(1 for r in rows if r["pal_landed"] is not None)
    tom_l = sum(1 for r in rows if r["tom_landed"] is not None)
    sell = sum(1 for r in rows if r["apbs_sell"] is not None)
    print(
        f"wrote {OUT_XLSX.name}: {n} rows "
        f"(Palconn FOB {pal}/{pal_l} landed, Tommur FOB {tom}/{tom_l} landed, Zhenpeng FOB {zhen}, site {sell})"
    )
    for needle in ("PEX-ELBOW", "PEX-B PIPE", "PVC-PIPE-SOLID", "CPVC-PIPE-SCH80", "CPVC-PIPE-SDR11", "CPVC-CPLNG"):
        hits = [r for r in rows if r["code"] == needle]
        for r in hits[:8]:
            print(
                f"  {r['item'][:56]:56} {r['size']:14} "
                f"P {r['pal_fob']} T {r['tom_fob']} Z {r['zhen_fob']} | "
                f"Pl {r['pal_landed']} Tl {r['tom_landed']} Zl {r['zhen_landed']} | sell {r['apbs_sell']}"
            )


if __name__ == "__main__":
    main()
